import argparse
import copy
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches


SPEC = importlib.util.spec_from_file_location("build_deck", Path(__file__).parents[1] / "scripts" / "build_deck.py")
build_deck = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(build_deck)


class BuildDeckTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)

    def build(self, deck, sections="steps", max_per_slide=5):
        args = argparse.Namespace(base=None, only=None, insert_at=None, sections=sections,
                                  max_per_slide=max_per_slide)
        builder = build_deck.Builder(deck, self.base, args)
        out = self.base / "test.pptx"
        builder.build(out)
        self.assertEqual(builder.overflows, [])
        return builder, Presentation(out)

    def test_unassigned_peek_appears_once_on_last_data_slide_before_operations(self):
        Image.new("RGB", (1280, 800), "white").save(self.base / "main.png")
        Image.new("RGB", (240, 70), "navy").save(self.base / "peek.png")
        deck = {
            "title": "取得と操作",
            "data_layout": [{"no": i, "name": name, "type": "文字列"}
                            for i, name in ((1, "商品名"), (2, "価格"))],
            "shots": [{
                "id": "mixed", "screen": "商品詳細", "image": "main.png",
                "steps": [
                    {"kind": "data", "item": 1, "box": [80, 80, 240, 40]},
                    {"kind": "data", "item": 2, "box": [80, 160, 240, 40]},
                    {"kind": "op", "text": "次へ進む", "box": [80, 260, 180, 40]},
                ],
                "peeks": [{"text": "価格が表示されない場合は空欄にします", "image": "peek.png"}],
            }],
        }
        _, prs = self.build(deck, max_per_slide=1)
        matches = [(i, shape) for i, slide in enumerate(prs.slides) for shape in slide.shapes
                   if shape.has_text_frame and shape.text == deck["shots"][0]["peeks"][0]["text"]]
        self.assertEqual([i for i, _ in matches], [1])
        self.assertEqual(prs.slides[1].shapes.title.text, "価格を取得（2/2）")
        pictures = [(i, shape) for i, slide in enumerate(prs.slides) for shape in slide.shapes
                    if shape.name == "表示差分の例"]
        self.assertEqual([i for i, _ in pictures], [1])
        self.assertEqual(pictures[0][1].image.blob, (self.base / "peek.png").read_bytes())
        self.assertEqual(len(prs.slides), 3)

    def test_summary_measures_off_screen_origins_and_keeps_notes_above_footer(self):
        origins = ["外部システムから取得する商品ごとの登録情報を集計した結果です", "利用者が別途管理する社内台帳から取得する分類情報と識別番号です"]
        self.assertTrue(all(29 <= len(origin) <= 31 for origin in origins))
        items = [{"no": i, "name": f"項目{i}", "type": "文字列"} for i in range(1, 14)]
        for item, origin in zip(items[-2:], origins):
            item.update(on_screen=False, origin=origin)
        builder, prs = self.build({"title": "取得元一覧", "shots": [], "data_layout": items}, "summary")
        self.assertEqual(len(prs.slides), 2)
        actual_names, actual_origins = [], []
        for slide in prs.slides:
            table = next(shape for shape in slide.shapes if shape.name == "取得項目一覧")
            footer = next(shape for shape in slide.shapes if shape.name == "フッター")
            self.assertLessEqual(table.top + table.height, Inches(builder.geo.body_b))
            for row in list(table.table.rows)[1:]:
                actual_names.append(row.cells[1].text)
                if row.cells[3].text in origins:
                    actual_origins.append(row.cells[3].text)
            for note in (shape for shape in slide.shapes if shape.name == "グレーアウト注記"):
                self.assertGreaterEqual(note.top, table.top + table.height)
                self.assertLessEqual(note.top + note.height, Inches(builder.geo.body_b))
                self.assertLess(note.top + note.height, footer.top)
        self.assertEqual(actual_names, [item["name"] for item in items])
        self.assertEqual(actual_origins, origins)

    def test_summary_reports_a_row_that_cannot_fit_on_one_page(self):
        deck = {"title": "長い取得元", "shots": [], "data_layout": [
            {"no": 1, "name": "外部項目", "type": "文字列", "on_screen": False,
             "origin": "取得元の詳しい説明" * 100},
        ]}
        args = argparse.Namespace(base=None, only=None, insert_at=None, sections="summary", max_per_slide=5)
        builder = build_deck.Builder(deck, self.base, args)
        out = self.base / "overflow.pptx"
        builder.build(out)
        self.assertEqual(len(builder.overflows), 1)
        self.assertIn("No.1", builder.overflows[0])
        self.assertIn("1ページに収まりません", builder.overflows[0])
        prs = Presentation(out)
        self.assertEqual(len(prs.slides), 1)
        table = next(shape for shape in prs.slides[0].shapes if shape.name == "取得項目一覧")
        self.assertGreater(table.top + table.height, Inches(builder.geo.body_b))
        self.assertEqual(table.table.cell(1, 3).text, deck["data_layout"][0]["origin"])
        scenario = self.base / "overflow.json"
        scenario.write_text(json.dumps(deck), encoding="utf-8")
        result = subprocess.run([sys.executable, str(Path(build_deck.__file__)), str(scenario),
                                 "--sections", "summary", "--out", str(out)], capture_output=True, check=False)
        self.assertEqual(result.returncode, 3, result.stderr)
        self.assertIn(b"NG:", result.stderr)


class MergeCapturesTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.captures = self.base / "captures"
        self.captures.mkdir()
        self.deck = {
            "shots": [{
                "id": "detail", "steps": [{"kind": "data", "item": 1, "target": "#title"}],
                "peeks": [
                    {"goto": "https://example.com/a", "target": "#a", "mark": "#missing-a", "text": "説明A"},
                    {"goto": "https://example.com/b", "target": "#b", "mark": "#missing-b", "text": "説明B"},
                ],
            }],
        }
        self.captured = copy.deepcopy(self.deck)
        shot = self.captured["shots"][0]
        shot.update(image="main.png", image_size=[1280, 800], scale=1)
        shot["steps"][0]["box"] = [80, 80, 240, 40]
        for i, peek in enumerate(shot["peeks"]):
            peek.update(image=f"peek-{i}.png", image_size=[240, 70], scale=1,
                        mark_box=[10, 10, 60, 20], captured_url=peek["goto"])
        self.write_captures()

    def write_captures(self):
        (self.captures / "deck.json").write_text(json.dumps(self.captured), encoding="utf-8")

    def merge(self, deck):
        return build_deck.merge_captures(deck, self.base, self.captures)

    def test_peek_reorder_requires_recapture(self):
        self.deck["shots"][0]["peeks"].reverse()
        with self.assertRaisesRegex(SystemExit, "peeks.*撮り直"):
            self.merge(self.deck)

    def test_each_changed_capture_parameter_requires_recapture(self):
        changes = {"goto": "https://example.com/other", "target": ["#a", "#extra"],
                   "mark": "#different", "pad": 20, "wait_ms": 3000}
        for key, value in changes.items():
            with self.subTest(key=key):
                deck = copy.deepcopy(self.deck)
                deck["shots"][0]["peeks"][0][key] = value
                with self.assertRaisesRegex(SystemExit, "peeks.*撮り直"):
                    self.merge(deck)
        del self.deck["shots"][0]["peeks"][0]["mark"]
        with self.assertRaisesRegex(SystemExit, "peeks.*撮り直"):
            self.merge(self.deck)

    def test_peek_wording_changes_keep_corresponding_images_and_coordinates(self):
        self.deck["shots"][0]["peeks"][0]["text"] = "修正した説明A"
        shot = self.merge(self.deck)["shots"][0]
        self.assertEqual(shot["peeks"][0]["text"], "修正した説明A")
        for i, peek in enumerate(shot["peeks"]):
            self.assertEqual(peek["image"], str((self.captures / f"peek-{i}.png").resolve()))
            self.assertEqual(peek["mark_box"], [10, 10, 60, 20])
        self.assertEqual(shot["steps"][0]["box"], [80, 80, 240, 40])

    def test_uncaptured_peek_requires_recapture(self):
        self.deck["shots"][0]["peeks"].append({"goto": "https://example.com/c", "target": "#c", "text": "説明C"})
        with self.assertRaisesRegex(SystemExit, "peeks.*撮り直"):
            self.merge(self.deck)

    def test_missing_peek_image_requires_recapture(self):
        del self.captured["shots"][0]["peeks"][0]["image"]
        self.write_captures()
        with self.assertRaisesRegex(SystemExit, "peeks.*撮り直"):
            self.merge(self.deck)

    def test_explicit_peek_image_uses_scenario_directory(self):
        peek = self.deck["shots"][0]["peeks"][0]
        peek.update(image="replacement.png", target="#replacement")
        shot = self.merge(self.deck)["shots"][0]
        self.assertEqual(shot["peeks"][0]["image"], str((self.base / "replacement.png").resolve()))


if __name__ == "__main__":
    unittest.main()
