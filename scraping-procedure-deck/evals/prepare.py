import argparse
import copy
import hashlib
import json
import os
import shutil
import subprocess
import sys
import threading
import urllib.error
import urllib.request
import uuid
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

EVALS = Path(__file__).resolve().parent
SKILL = EVALS.parent
FILES = EVALS / "files"
SITE = FILES / "mock-shop" / "site"
GEN = FILES / "generated"
PORT = 8765
BASE_URL = f"http://localhost:{PORT}/"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
SECTION_EXT = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"


class QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass


def site_root(rev):
    if rev == 1:
        return SITE
    root = GEN / f"site-rev{rev}"
    shutil.rmtree(root, ignore_errors=True)
    shutil.copytree(SITE, root)
    (root / "rev.js").write_text(f"window.SITE_REV = {rev};\n", encoding="utf-8")
    return root


def serve(rev):
    root = site_root(rev)
    print(f"{BASE_URL} -> {root}（Ctrl+C で停止）", flush=True)
    ThreadingHTTPServer(("127.0.0.1", PORT), partial(QuietHandler, directory=str(root))).serve_forever()


def ensure_server():
    try:
        with urllib.request.urlopen(BASE_URL + "index.html", timeout=2):
            pass
    except urllib.error.HTTPError as exc:
        raise RuntimeError(f"既存サーバー {BASE_URL} の index.html が HTTP {exc.code} です。正規 rev1 を確認できないため中止します。") from exc
    except OSError:
        try:
            server = ThreadingHTTPServer(("127.0.0.1", PORT), partial(QuietHandler, directory=str(SITE)))
        except OSError as bind_error:
            raise RuntimeError(f"{BASE_URL} の正規 rev1 を確認できず、既定サイトの起動もできません。既存サービスは停止していません。") from bind_error
        threading.Thread(target=server.serve_forever, daemon=True).start()
        return
    try:
        with urllib.request.urlopen(BASE_URL + "rev.js", timeout=2) as response:
            revision = response.read()
    except OSError as exc:
        raise RuntimeError(f"既存サーバー {BASE_URL} の rev.js を取得できません。別サイトの可能性があるため中止します。") from exc
    if revision.strip() != (SITE / "rev.js").read_bytes().strip():
        raise RuntimeError(f"既存サーバー {BASE_URL} の rev.js が正規 rev1 と一致しません。改修版 rev2 または別サイトを採用せず中止します。既存サービスは停止していません。")


def run(cmd):
    print("$", " ".join(str(c) for c in cmd), flush=True)
    subprocess.run([str(c) for c in cmd], check=True, env={**os.environ, "PYTHONIOENCODING": "utf-8"})


# ---------- 既存資料（差し込み先）のテンプレート ----------

def edit_theme(prs, fonts, colors, name):
    part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    root.set("name", name)
    for kind, (latin, ea) in fonts.items():
        f = root.find(f".//{A}fontScheme/{A}{kind}Font")
        f.find(f"{A}latin").set("typeface", latin)
        f.find(f"{A}ea").set("typeface", ea)
        for alt in f.findall(f"{A}font"):
            if alt.get("script") == "Jpan":
                if ea:
                    alt.set("typeface", ea)
                else:
                    f.remove(alt)
    scheme = root.find(f".//{A}clrScheme")
    for key, val in colors.items():
        el = scheme.find(f"{A}{key}")
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, f"{A}srgbClr", val=val)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def rescale(prs, width, height):
    sx, sy = width / prs.slide_width, height / prs.slide_height
    master = prs.slide_masters[0]
    for el in [master._element] + [lay._element for lay in master.slide_layouts]:
        for xfrm in el.iter(qn("a:xfrm")):
            off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
            if off is not None:
                off.set("x", str(int(int(off.get("x")) * sx)))
                off.set("y", str(int(int(off.get("y")) * sy)))
            if ext is not None:
                ext.set("cx", str(int(int(ext.get("cx")) * sx)))
                ext.set("cy", str(int(int(ext.get("cy")) * sy)))
    prs.slide_width, prs.slide_height = Emu(int(width)), Emu(int(height))


def title_style(prs, size_pt, bold, align=None):
    ppr = prs.slide_masters[0]._element.find(f"{qn('p:txStyles')}/{qn('p:titleStyle')}/{A}lvl1pPr")
    rpr = ppr.find(f"{A}defRPr")
    rpr.set("sz", str(size_pt * 100))
    rpr.set("b", "1" if bold else "0")
    if align:
        ppr.set("algn", align)


def place_master(prs, kind, rect):
    for ph in prs.slide_masters[0].placeholders:
        if ph.placeholder_format.type == kind:
            ph.left, ph.top, ph.width, ph.height = (Inches(v) for v in rect)


def drop_placeholders(prs, kinds):
    master = prs.slide_masters[0]
    for owner in [master] + list(master.slide_layouts):
        for ph in list(owner.placeholders):
            if ph.placeholder_format.type in kinds:
                ph._element.getparent().remove(ph._element)


def master_decoration(prs, shapes):
    # レイアウト共通の飾り（タイトル下の線・ロゴ文字）を、一時スライドで作った図形をマスターへ移して付ける
    tmp = prs.slides.add_slide(prs.slide_layouts[6])
    tree = prs.slide_masters[0]._element.find(qn("p:cSld")).find(qn("p:spTree"))
    next_id = max(int(e.get("id")) for e in tree.iter(qn("p:cNvPr"))) + 1
    for make in shapes:
        shp = make(tmp.shapes)
        el = shp._element
        el.getparent().remove(el)
        el.find(".//" + qn("p:cNvPr")).set("id", str(next_id))
        next_id += 1
        tree.append(el)
    sld_ids = prs.slides._sldIdLst
    rid = sld_ids[-1].rId
    prs.part.drop_rel(rid)
    sld_ids.remove(sld_ids[-1])


def add_footers(slide, text):
    for ph in slide.slide_layout.placeholders:
        kind = ph.placeholder_format.type
        if kind not in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER):
            continue
        el = copy.deepcopy(ph._element)
        el.find(".//" + qn("p:cNvPr")).set("id", str(slide.shapes._next_shape_id))
        slide.shapes._spTree.append(el)
        if kind == PP_PLACEHOLDER.FOOTER:
            body = el.find(qn("p:txBody"))
            for p in body.findall(qn("a:p")):
                body.remove(p)
            p = etree.SubElement(body, qn("a:p"))
            r = etree.SubElement(p, qn("a:r"))
            etree.SubElement(r, qn("a:rPr"), lang="ja-JP", altLang="en-US")
            etree.SubElement(r, qn("a:t")).text = text


def add_sections(prs, groups):
    ids = [s.slide_id for s in prs.slides]
    root = prs.part._element
    ext_lst = root.find(qn("p:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(root, qn("p:extLst"))
    ext = etree.SubElement(ext_lst, qn("p:ext"), uri=SECTION_EXT)
    lst = etree.SubElement(ext, f"{{{P14}}}sectionLst", nsmap={"p14": P14})
    for name, idxs in groups:
        sec = etree.SubElement(lst, f"{{{P14}}}section", name=name,
                               id="{" + str(uuid.uuid5(uuid.NAMESPACE_URL, name)).upper() + "}")
        sl = etree.SubElement(sec, f"{{{P14}}}sldIdLst")
        for i in idxs:
            etree.SubElement(sl, f"{{{P14}}}sldId", id=str(ids[i]))


def layout(prs, name):
    return next(lay for lay in prs.slide_layouts if lay.name == name)


def fill_deck(prs, cover, pages, footer):
    s = prs.slides.add_slide(layout(prs, "Title Slide"))
    s.shapes.title.text = cover[0]
    s.placeholders[1].text = cover[1]
    for title, bullets in pages:
        s = prs.slides.add_slide(layout(prs, "Title and Content"))
        s.shapes.title.text = title
        body = s.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b
        if footer:
            add_footers(s, footer)


PAGES = [
    ("ご提案の背景", ["競合サイトの価格・在庫を毎日人手で確認しており、担当者の負荷が高い", "確認漏れによる価格改定の遅れが発生している"]),
    ("収集対象と取得項目", ["対象サイト：EC モール 3 サイト", "取得項目：商品名・価格・ポイント・在庫状況など 12 項目", "取得頻度：1 日 1 回（深夜帯）"]),
    ("スケジュール", ["10 月：取得手順の確定・テスト収集", "11 月：本番収集開始・納品データの確認", "12 月：運用定着・改善"]),
    ("お見積り", ["初期費用：一式", "月額費用：サイト数 × 取得頻度に応じて算出"]),
]


def make_templates(out):
    out.mkdir(parents=True, exist_ok=True)
    made = {}

    prs = Presentation()
    rescale(prs, Inches(13.333), Inches(7.5))
    edit_theme(prs, {"major": ("游ゴシック", "游ゴシック"), "minor": ("游ゴシック", "游ゴシック")},
               {"dk2": "1F3864", "accent1": "1F4E79", "accent2": "C55A11", "hlink": "0563C1"}, "サンプル社 提案書")
    title_style(prs, 28, True, "l")
    place_master(prs, PP_PLACEHOLDER.TITLE, (0.5, 0.3, 12.33, 0.85))
    place_master(prs, PP_PLACEHOLDER.BODY, (0.5, 1.45, 12.33, 5.2))
    master_decoration(prs, [
        lambda sh: _bar(sh, Inches(0.5), Inches(1.2), Inches(12.33), Inches(0.05), "1F4E79"),
        lambda sh: _logo(sh, "SAMPLE Inc.", Inches(11.0), Inches(0.05), Inches(2.0), Inches(0.3)),
    ])
    fill_deck(prs, ("データ収集サービスのご提案", "株式会社サンプル　2026年10月"), PAGES, "株式会社サンプル　社外秘")
    add_sections(prs, [("はじめに", [0, 1]), ("収集内容", [2]), ("進め方", [3, 4])])
    made["proposal_16x9.pptx"] = prs

    prs = Presentation()
    edit_theme(prs, {"major": ("ＭＳ Ｐゴシック", "ＭＳ Ｐゴシック"), "minor": ("ＭＳ Ｐゴシック", "ＭＳ Ｐゴシック")},
               {"accent1": "2E7D32", "dk2": "263238"}, "旧テンプレート")
    title_style(prs, 32, True)
    prs.slide_layouts.remove(layout(prs, "Title Only"))
    fill_deck(prs, ("データ収集 業務手順書", "情報システム部"), PAGES[:3], "社外秘")
    made["legacy_4x3.pptx"] = prs

    prs = Presentation()
    rescale(prs, Inches(10), Inches(5.625))
    edit_theme(prs, {"major": ("Arial", ""), "minor": ("Arial", "")},
               {"accent1": "4285F4", "dk2": "202124"}, "Simple Light")
    title_style(prs, 30, False)
    drop_placeholders(prs, (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER))
    for lay, new in (("Title Slide", "TITLE"), ("Title and Content", "TITLE_AND_BODY"), ("Title Only", "TITLE_ONLY"),
                     ("Section Header", "SECTION_HEADER"), ("Blank", "BLANK")):
        layout(prs, lay).name = new
    s = prs.slides.add_slide(layout(prs, "TITLE"))
    s.shapes.title.text = "クローラー導入のご提案"
    s.placeholders[1].text = "Google スライドで作成"
    for title, bullets in PAGES[:2]:
        s = prs.slides.add_slide(layout(prs, "TITLE_AND_BODY"))
        s.shapes.title.text = title
        s.placeholders[1].text_frame.text = "\n".join(bullets)
        add_footers(s, None)
    made["gslides_like.pptx"] = prs

    prs = Presentation()
    rescale(prs, Inches(13.333), Inches(7.5))
    edit_theme(prs, {"major": ("BIZ UDPゴシック", "BIZ UDPゴシック"), "minor": ("BIZ UDPゴシック", "BIZ UDPゴシック")},
               {"dk1": "0F172A", "lt1": "F8FAFC", "dk2": "1E293B", "lt2": "CBD5E1", "accent1": "38BDF8",
                "accent2": "F472B6", "hlink": "7DD3FC"}, "Dark")
    title_style(prs, 30, True)
    cmap = prs.slide_masters[0]._element.find(qn("p:clrMap"))
    cmap.set("bg1", "dk1")
    cmap.set("tx1", "lt1")
    cmap.set("bg2", "dk2")
    cmap.set("tx2", "lt2")
    fill_deck(prs, ("スクレイピング運用レポート", "2026年10月"), PAGES[1:3], "Confidential")
    made["dark_16x9.pptx"] = prs

    hashes = {}
    for name, p in made.items():
        path = out / name
        p.save(path)
        hashes[name] = hashlib.sha256(path.read_bytes()).hexdigest()
        print(f"{path}  {len(p.slides)} slides  {Emu(p.slide_width).inches:.3f}x{Emu(p.slide_height).inches:.3f}in")
    (out / "hashes.json").write_text(json.dumps(hashes, indent=2), encoding="utf-8")


def _bar(shapes, x, y, w, h, rgb):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    shp = shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(rgb)
    shp.line.fill.background()
    shp.name = "タイトル下線"
    return shp


def _logo(shapes, text, x, y, w, h):
    tb = shapes.add_textbox(x, y, w, h)
    tb.text_frame.text = text
    run = tb.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(12)
    run.font.bold = True
    tb.name = "ロゴ"
    return tb


# ---------- モックサイトの撮影済み作業フォルダ（差し込み・修正・改修シナリオ用） ----------

def make_mock_work(out, edited_dir):
    ensure_server()
    shutil.rmtree(out, ignore_errors=True)
    (out / "out").mkdir(parents=True)
    shutil.copy(FILES / "mock-shop" / "scenario.json", out / "scenario.json")
    run([sys.executable, SKILL / "scripts" / "capture.py", out / "scenario.json", "--out", out / "captures"])
    deck = out / "out" / "zbmart.pptx"
    run([sys.executable, SKILL / "scripts" / "build_deck.py", out / "scenario.json", "--captures", out / "captures",
         "-o", deck])
    edited_dir.mkdir(parents=True, exist_ok=True)
    hand_edit(deck, edited_dir / "zbmart_手直し済み.pptx", edited_dir / "hand_edits.json")


def hand_edit(src, dst, record):
    # PowerPoint 上での手直しを模す: 赤枠を少し動かし、タイトルを書き換え、注記を足す
    prs = Presentation(src)
    slide = next(s for s in prs.slides if s.shapes.title is not None and "（1/3）" in s.shapes.title.text_frame.text)
    frame = next(sh for sh in slide.shapes if sh.name.startswith("赤枠④"))
    frame.left, frame.top = frame.left + Inches(0.06), frame.top - Inches(0.03)
    slide.shapes.title.text_frame.paragraphs[0].runs[0].text = "商品の基本情報を取得（1/3）"
    memo = slide.shapes.add_textbox(Inches(5.6), Inches(1.1), Inches(3.4), Inches(0.28))
    memo.text_frame.text = "※価格は2026年10月3日時点の表示"
    memo.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    memo.name = "手直しメモ"
    prs.save(dst)
    record.write_text(json.dumps({
        "file": dst.name,
        "slide_title": "商品の基本情報を取得（1/3）",
        "frame": {"name": frame.name, "left": int(frame.left), "top": int(frame.top)},
        "memo": {"name": "手直しメモ", "text": "※価格は2026年10月3日時点の表示"},
        "slides": len(prs.slides),
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{dst}  {len(prs.slides)} slides（手直し内容: {record.name}）")


# ---------- C（手持ち画像）用の画像と正解座標 ----------

SHOTS = [
    {"name": "01_検索結果.png", "url": "search.html?q=%E3%83%AF%E3%82%A4%E3%83%A4%E3%83%AC%E3%82%B9%E3%82%A4%E3%83%A4%E3%83%9B%E3%83%B3&shop=zb-denki",
     "viewport": [1366, 768], "scale": 1.5,
     "boxes": {"op:shop": ("aside section:has(h4:text-is('ショップで絞り込み')) a:text-is('ZBデンキ')", "box"),
               "op:chip": ("[class*='ConditionChip_active']", "box"),
               "op:first": ("[class*='ItemCard_title'] >> nth=0", "box")}},
    {"name": "02_商品詳細.png", "url": "item.html?id=1001", "viewport": [1366, 768], "scale": 1.5,
     "boxes": {"data:3": ("h1[class*='ItemTitle_itemTitle']", "text"), "data:4": ("strong[class*='Price_current']", "text"),
               "data:5": ("[class*='Point_point']", "text"), "data:6": ("a[class*='Review_count']", "text"),
               "data:7": ("[class*='Review_score']", "text"), "data:11": ("[class*='Stock_row']", "text")}},
    {"name": "03_商品仕様.png", "url": "item.html?id=1001", "viewport": [1366, 768], "scale": 1.5, "banner": True,
     "scroll_center": ("[class*='SpecTable_table'] tr:has(th:text-is('JANコード'))", 728),
     "boxes": {"data:8": ("[class*='SpecTable_table'] tr:has(th:text-is('メーカー'))", "text"),
               "data:9": ("[class*='SpecTable_table'] tr:has(th:text-is('型番'))", "text"),
               "data:10": ("[class*='SpecTable_table'] tr:has(th:text-is('JANコード'))", "text")}},
    {"name": "会員ログイン後_商品詳細.png", "url": "item.html?id=1001", "viewport": [1280, 720], "scale": 1.25, "member": True,
     "boxes": {"data:3": ("h1[class*='ItemTitle_itemTitle']", "text"), "data:13": ("[class*='MemberPrice_member']", "text")}},
]


def make_screens(out):
    sys.path.insert(0, str(SKILL / "scripts"))
    from capture import RECT_JS
    from playwright.sync_api import sync_playwright

    ensure_server()
    out.mkdir(parents=True, exist_ok=True)
    truth = {}
    with sync_playwright() as p:
        browser = p.chromium.launch()
        for spec in SHOTS:
            vw, vh = spec["viewport"]
            ctx = browser.new_context(viewport={"width": vw, "height": vh}, device_scale_factor=spec["scale"],
                                      locale="ja-JP", timezone_id="Asia/Tokyo")
            if spec.get("member"):
                ctx.add_init_script("sessionStorage.setItem('zb-member', '1')")
            page = ctx.new_page()
            page.goto(BASE_URL + spec["url"])
            page.wait_for_timeout(2500)
            if not spec.get("banner"):
                page.evaluate("() => document.querySelectorAll(\"[class*='CookieBanner_banner']\").forEach(e => e.remove())")
            if spec.get("scroll_center"):
                sel, y = spec["scroll_center"]
                r = page.locator(sel).first.evaluate(RECT_JS, "box")
                page.evaluate("y => window.scrollTo(0, y)", r[1] + r[3] / 2 - y)
                page.wait_for_timeout(600)
            page.screenshot(path=str(out / spec["name"]))
            sy = page.evaluate("() => window.scrollY")
            boxes = {}
            for key, (sel, fit) in spec["boxes"].items():
                x, y, w, h = page.locator(sel).first.evaluate(RECT_JS, fit)
                s = spec["scale"]
                boxes[key] = [round(x * s, 1), round((y - sy) * s, 1), round(w * s, 1), round(h * s, 1)]
            truth[spec["name"]] = {"scale": spec["scale"], "size": [int(vw * spec["scale"]), int(vh * spec["scale"])],
                                   "url": BASE_URL + spec["url"], "boxes": boxes}
            print(out / spec["name"])
            ctx.close()
        browser.close()
    (out / "truth.json").write_text(json.dumps(truth, ensure_ascii=False, indent=2), encoding="utf-8")


# ---------- データレイアウト（Excel）とPlaywright無し環境 ----------

def make_jma_layout(path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "データレイアウト"
    ws["A1"] = "気象データ（日別）納品レイアウト　Ver.1.2"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:F1")
    header = ["項目No", "項目名", "データ型", "桁数", "画面表示", "備考"]
    rows = [
        [1, "取得日時", "日時", "", "なし", "クローラーの実行時刻"],
        [2, "地点名", "文字列", 20, "あり", "例: 東京"],
        [3, "年月日", "日付", "", "あり", "見出しの年月と表の「日」を組み合わせる"],
        [4, "平均気温", "数値", "5,1", "あり", "単位 ℃"],
        [5, "最高気温", "数値", "5,1", "あり", "単位 ℃"],
        [6, "最低気温", "数値", "5,1", "あり", "単位 ℃"],
        [7, "降水量合計", "数値", "6,1", "あり", "単位 mm。「--」は0とする"],
        [8, "日照時間", "数値", "4,1", "あり", "単位 h"],
        [9, "取得元URL", "文字列", "", "なし", "表示したページのURL"],
    ]
    for c, v in enumerate(header, 1):
        cell = ws.cell(row=3, column=c, value=v)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows, 4):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    for col, w in zip("ABCDEF", (8, 14, 10, 8, 10, 40)):
        ws.column_dimensions[col].width = w
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(path)


def make_envs(out):
    venv = out / "venv-noplaywright"
    if not venv.exists():
        run([sys.executable, "-m", "venv", venv])
        py = venv / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
        run([py, "-m", "pip", "install", "-q", "python-pptx", "Pillow", "lxml", "openpyxl"])
    (out / "empty-browsers").mkdir(parents=True, exist_ok=True)
    print(f"Playwright無しの Python: {venv}")
    print(f"ブラウザ未導入を再現する PLAYWRIGHT_BROWSERS_PATH: {out / 'empty-browsers'}")


def main():
    ap = argparse.ArgumentParser(description="scraping-procedure-deck の検証シナリオ用フィクスチャを作る")
    sub = ap.add_subparsers(dest="cmd", required=True)
    sv = sub.add_parser("serve", help="モックサイトを http://localhost:8765/ で配信する（前面で動き続ける）")
    sv.add_argument("--rev", type=int, default=1, help="2 でサイト改修後の版を配信する")
    sub.add_parser("all", help="templates / mock / screens / layouts をまとめて作る")
    sub.add_parser("templates", help="差し込み先の既存資料4種")
    sub.add_parser("mock", help="モックサイトを撮影済みの作業フォルダと、手直し済み資料")
    sub.add_parser("screens", help="C 用の手持ち画像と正解座標")
    sub.add_parser("layouts", help="Excel のデータレイアウト")
    sub.add_parser("envs", help="Playwright 無しの venv と、空のブラウザ置き場")
    args = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    if args.cmd == "serve":
        serve(args.rev)
        return
    steps = ["templates", "mock", "screens", "layouts"] if args.cmd == "all" else [args.cmd]
    for step in steps:
        print(f"== {step}", flush=True)
        if step == "templates":
            make_templates(GEN / "templates")
        elif step == "mock":
            make_mock_work(GEN / "mock-work", GEN / "hand-edited")
        elif step == "screens":
            make_screens(GEN / "c-screens")
        elif step == "layouts":
            make_jma_layout(GEN / "jma" / "気象データ_データレイアウト.xlsx")
        elif step == "envs":
            make_envs(GEN)


if __name__ == "__main__":
    main()
