import argparse
import colorsys
import hashlib
import json
import math
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu

P14 = "{http://schemas.microsoft.com/office/powerpoint/2010/main}"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
FOOTERS = (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE)
COVER_LAYOUTS = ("Title Slide", "タイトル スライド", "TITLE")
TECH = [
    (r"\[class|class\*=", "CSSの属性セレクタ"),
    (r":has\(|:text-is\(|:nth-child|>>\s*nth=", "セレクタの疑似クラス"),
    (r"get_by_\w+|getBy\w+|locator\(|query_selector", "Playwrightのロケータ"),
    (r"xpath|(?<![\w:/])//[a-z]+", "XPath"),
    (r"data-testid|test_id", "テスト用ID"),
    (r"\b[A-Z][A-Za-z]+_[A-Za-z]+__[A-Za-z0-9]{4,}", "ハッシュ付きクラス名"),
    (r"(?<![\w&])#[a-z][\w-]{2,}", "CSSのID（要確認）"),
]
URL = re.compile(r"https?://[^\s）)」]+")
MARKER = re.compile(r"^[\s①-⑳㉑-㉟㊱-㊿※0-9()]+$")
# 画面から取らない項目のグレーアウトは意図した低コントラストなので対象外にする
BY_DESIGN = {("#7F7F7F", "#D9D9D9")}


def shapes_of(container):
    for shp in container:
        if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from shapes_of(shp.shapes)
        else:
            yield shp


def frames_of(shp):
    if shp.has_text_frame:
        yield shp.text_frame
    if getattr(shp, "has_table", False) and shp.has_table:
        for row in shp.table.rows:
            for cell in row.cells:
                yield cell.text_frame


def texts(slide):
    out = [(shp.name, tf.text) for shp in shapes_of(slide.shapes) for tf in frames_of(shp)]
    if slide.has_notes_slide:
        out.append(("ノート", slide.notes_slide.notes_text_frame.text))
    return out


def links(slide):
    found = []
    for shp in shapes_of(slide.shapes):
        for tf in frames_of(shp):
            for p in tf.paragraphs:
                for r in p.runs:
                    if r.hyperlink.address:
                        found.append((r.text, r.hyperlink.address))
    return found


def title_of(slide):
    t = slide.shapes.title
    return t.text_frame.text if t is not None else ""


def named(slide, prefix):
    return [shp for shp in shapes_of(slide.shapes) if shp.name.startswith(prefix)]


def sections(prs):
    ext = prs.part._element.find(qn("p:extLst"))
    if ext is None:
        return None
    lst = ext.find(f".//{P14}sectionLst")
    if lst is None:
        return None
    return [(sec.get("name"), [int(x.get("id")) for x in sec.iter(f"{P14}sldId")]) for sec in lst.findall(f"{P14}section")]


def summary(prs):
    secs = sections(prs)
    where = {sid: name for name, ids in (secs or []) for sid in ids}
    slides = []
    for i, s in enumerate(prs.slides, 1):
        badge = next((shp.text_frame.text for shp in named(s, "区分ラベル")), "")
        screen = next((shp.text_frame.text for shp in named(s, "画面名")), "")
        slides.append({
            "no": i, "layout": s.slide_layout.name, "section": where.get(s.slide_id), "title": title_of(s),
            "kind": badge, "screen": screen,
            "frames": [shp.name for shp in named(s, "赤枠")], "labels": [shp.text_frame.text for shp in named(s, "番号")],
            "pictures": sum(1 for shp in shapes_of(s.shapes) if shp.shape_type == MSO_SHAPE_TYPE.PICTURE),
            "links": links(s),
            "footer": any(ph.placeholder_format.type in FOOTERS for ph in s.placeholders) or bool(named(s, "フッター")),
        })
    return {"size_in": [round(Emu(prs.slide_width).inches, 3), round(Emu(prs.slide_height).inches, 3)],
            "slides": slides, "sections": secs}


def cmd_summary(args):
    data = summary(Presentation(args.pptx))
    if args.json:
        print(json.dumps(data, ensure_ascii=False, indent=2))
        return 0
    print(f"サイズ {data['size_in'][0]}x{data['size_in'][1]}in  {len(data['slides'])}枚")
    for s in data["slides"]:
        sec = f"［{s['section']}］" if s["section"] else ""
        kind = f"〈{s['kind']}〉" if s["kind"] else ""
        print(f"{s['no']:>2}. {sec}{kind}{s['title']}  | 画面: {s['screen'] or '-'} | 枠{len(s['frames'])} 番号{''.join(s['labels'])}"
              f" | 画像{s['pictures']} | リンク{len(s['links'])} | フッター{'あり' if s['footer'] else 'なし'} | {s['layout']}")
    return 0


def cmd_leaks(args):
    prs = Presentation(args.pptx)
    bad = 0
    for i, s in enumerate(prs.slides, 1):
        cover = s.slide_layout.name in COVER_LAYOUTS
        for name, text in texts(s):
            for pat, label in TECH:
                for m in re.finditer(pat, text, re.I):
                    bad += 1
                    print(f"NG  {i}枚目 {name}: {label}「{m.group(0)}」 … {text[:60]!r}")
            for m in URL.finditer(text):
                if not cover:
                    bad += 1
                    print(f"NG  {i}枚目 {name}: URLが文字として表示されている「{m.group(0)}」")
    print("OK  技術情報・URLの表示なし" if not bad else f"計 {bad} 件")
    return 1 if bad else 0


class Checks:
    def __init__(self):
        self.failed = 0

    def __call__(self, label, cond, detail=""):
        self.failed += 0 if cond else 1
        print(f"{'OK ' if cond else 'NG '} {label}" + (f"  … {detail}" if detail else ""))


def signature(slide):
    return slide.slide_layout.name, title_of(slide), tuple(t for _, t in texts(slide))


def cmd_insert(args):
    out, base = Presentation(args.pptx), Presentation(args.base)
    ok = Checks()
    ok("スライドサイズが差し込み先と同じ", (out.slide_width, out.slide_height) == (base.slide_width, base.slide_height),
       f"{Emu(out.slide_width).inches:.3f}x{Emu(out.slide_height).inches:.3f}in")
    bs, os_ = [signature(s) for s in base.slides], [signature(s) for s in out.slides]
    n = len(os_) - len(bs)
    ok("スライドが追加されている", n > 0, f"{len(bs)} → {len(os_)} 枚")
    at = args.at
    placed = n > 0 and os_[:at - 1] == bs[:at - 1] and os_[at - 1 + n:] == bs[at - 1:]
    actual = next((k for k in range(1, len(bs) + 2) if os_[:k - 1] == bs[:k - 1] and os_[k - 1 + n:] == bs[k - 1:]), None)
    ok(f"元のスライドが順序どおり残り、追加分が {at} 枚目から連続している", placed,
       "" if placed else (f"実際は {actual} 枚目から" if actual else "元のスライドが変更・削除されている"))
    new = list(out.slides)[at - 1:at - 1 + n] if placed else []
    names = {lay.name for lay in base.slide_layouts}
    ok("追加スライドが差し込み先のマスター・レイアウトを使っている",
       new and len(out.slide_masters) == len(base.slide_masters) and all(s.slide_layout.name in names for s in new),
       ", ".join(sorted({s.slide_layout.name for s in new})))
    ok("表紙を追加していない", not any(s.slide_layout.name in COVER_LAYOUTS for s in new))
    empty = [at + k for k, s in enumerate(new) for ph in s.placeholders
             if ph.placeholder_format.type not in FOOTERS + (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)]
    ok("追加スライドに本文などの空プレースホルダーが残っていない", not empty, f"残っている: {empty}" if empty else "")
    base_footer = any(ph.placeholder_format.type in FOOTERS for s in base.slides for ph in s.placeholders)
    if base_footer:
        lacking = [at + k for k, s in enumerate(new) if not any(ph.placeholder_format.type in FOOTERS for ph in s.placeholders)]
        ok("追加スライドにも差し込み先のフッター／ページ番号がある", new and not lacking, f"なし: {lacking}" if lacking else "")
    secs = sections(out)
    if sections(base) is not None:
        ids_in = {sid for _, ids in secs or [] for sid in ids}
        ok("全スライドがいずれかのセクションに属している（修復ダイアログが出ない）",
           all(s.slide_id in ids_in for s in out.slides))
        if new and at > 1:
            prev = list(out.slides)[at - 2].slide_id
            sec = next((name for name, ids in secs if prev in ids), None)
            ok(f"追加スライドが直前のスライドと同じセクション「{sec}」に入っている",
               all(any(s.slide_id in ids for name, ids in secs if name == sec) for s in new))
    if args.hash:
        h = json.loads(Path(args.hash).read_text(encoding="utf-8")).get(Path(args.base).name)
        ok("差し込み先の元ファイルが書き換わっていない", hashlib.sha256(Path(args.base).read_bytes()).hexdigest() == h)
    return 1 if ok.failed else 0


def cmd_edits(args):
    rec = json.loads(Path(args.record).read_text(encoding="utf-8"))
    prs = Presentation(args.pptx)
    ok = Checks()
    slide = next((s for s in prs.slides if title_of(s) == rec["slide_title"]), None)
    ok(f"手直ししたタイトル「{rec['slide_title']}」が残っている", slide is not None)
    if slide is not None:
        frame = next((shp for shp in slide.shapes if shp.name == rec["frame"]["name"]), None)
        ok(f"手で動かした「{rec['frame']['name']}」の位置が保たれている",
           frame is not None and abs(frame.left - rec["frame"]["left"]) <= 1 and abs(frame.top - rec["frame"]["top"]) <= 1)
        memo = next((shp for shp in slide.shapes if shp.name == rec["memo"]["name"]), None)
        ok("手で足した注記が残っている", memo is not None and memo.text_frame.text == rec["memo"]["text"])
    ok("スライドが追加されている", len(prs.slides) > rec["slides"], f"{rec['slides']} → {len(prs.slides)} 枚")
    return 1 if ok.failed else 0


class Colors:
    def __init__(self, slide):
        master = slide.slide_layout.slide_master
        scheme = etree.fromstring(master.part.part_related_by(RT.THEME).blob).find(f".//{{{A_NS}}}clrScheme")
        self.theme = {}
        for el in scheme:
            c = el[0]
            self.theme[etree.QName(el).localname] = c.get("val") if etree.QName(c).localname == "srgbClr" else c.get("lastClr", "000000")
        self.map = dict(master._element.find(qn("p:clrMap")).attrib)
        for owner in (slide.slide_layout._element, slide._element):
            o = owner.find(f"{qn('p:clrMapOvr')}/{qn('a:overrideClrMapping')}")
            if o is not None:
                self.map = dict(o.attrib)
        self.background = self.slide_background(slide)

    def resolve(self, el):
        tag = etree.QName(el).localname
        if tag == "srgbClr":
            hexv = el.get("val")
        elif tag == "schemeClr":
            key = el.get("val")
            hexv = self.theme.get(self.map.get(key, key), "000000")
        elif tag == "sysClr":
            hexv = el.get("lastClr", "000000")
        else:
            hexv = {"white": "FFFFFF", "red": "FF0000"}.get(el.get("val"), "000000")
        rgb = [int(hexv[i:i + 2], 16) / 255 for i in (0, 2, 4)]
        mods = {etree.QName(m).localname: int(m.get("val")) / 100000 for m in el}
        if "lumMod" in mods or "lumOff" in mods:
            h, light, s = colorsys.rgb_to_hls(*rgb)
            light = min(1.0, max(0.0, light * mods.get("lumMod", 1) + mods.get("lumOff", 0)))
            rgb = list(colorsys.hls_to_rgb(h, light, s))
        return tuple(rgb)

    def fill(self, parent):
        sf = parent.find(qn("a:solidFill")) if parent is not None else None
        return self.resolve(sf[0]) if sf is not None and len(sf) else None

    def slide_background(self, slide):
        for owner in (slide._element, slide.slide_layout._element, slide.slide_layout.slide_master._element):
            bg = owner.find(f"{qn('p:cSld')}/{qn('p:bg')}")
            if bg is None:
                continue
            pr = bg.find(qn("p:bgPr"))
            if pr is not None and self.fill(pr):
                return self.fill(pr)
            ref = bg.find(qn("p:bgRef"))
            if ref is not None and len(ref):
                return self.resolve(ref[0])
        return self.scheme("bg1")

    def scheme(self, key):
        return self.resolve(etree.Element(f"{{{A_NS}}}schemeClr", val=key))


def luminance(rgb):
    def ch(c):
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = (ch(c) for c in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def ratio(a, b):
    la, lb = sorted((luminance(a), luminance(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def hexs(rgb):
    return "#" + "".join(f"{round(c * 255):02X}" for c in rgb)


def cmd_contrast(args):
    prs = Presentation(args.pptx)
    first, last = (int(v) for v in (args.slides or f"1-{len(prs.slides)}").split("-"))
    low = 0
    for i, s in enumerate(prs.slides, 1):
        if not first <= i <= last:
            continue
        col = Colors(s)
        default_text = col.scheme("tx1")
        filled = []
        for shp in shapes_of(s.shapes):
            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                filled.append((shp, None))
                continue
            own = col.fill(shp._element.find(qn("p:spPr")))
            cells = []
            if getattr(shp, "has_table", False) and shp.has_table:
                cells = [(c.text_frame, col.fill(c._tc.find(qn("a:tcPr"))) or own) for r in shp.table.rows for c in r.cells]
            elif shp.has_text_frame:
                center = (shp.left + shp.width / 2, shp.top + shp.height / 2)
                under = next((f for o, f in reversed(filled) if o.left <= center[0] <= o.left + o.width
                              and o.top <= center[1] <= o.top + o.height), col.background)
                cells = [(shp.text_frame, own or under)]
            if own is not None:
                filled.append((shp, own))
            for tf, bg in cells:
                if bg is None:
                    continue
                for p in tf.paragraphs:
                    for r in p.runs:
                        if not r.text.strip():
                            continue
                        rpr = r._r.find(qn("a:rPr"))
                        fg = col.fill(rpr) or default_text
                        size = r.font.size.pt if r.font.size else 18
                        large = size >= 18 or (size >= 14 and r.font.bold) or MARKER.match(r.text)
                        need = 3.0 if large else 4.5
                        cr = ratio(fg, bg)
                        if cr < need and (hexs(fg), hexs(bg)) not in BY_DESIGN:
                            low += 1
                            print(f"NG  {i}枚目 {shp.name}: 「{r.text[:20]}」 文字{hexs(fg)} / 背景{hexs(bg)}  コントラスト比 {cr:.2f}（基準 {need}）")
    print("OK  文字と背景のコントラスト比はすべて基準以上" if not low else f"計 {low} 件")
    return 1 if low else 0


def iou(a, b):
    w = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
    h = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    inter = w * h if w > 0 and h > 0 else 0.0
    return inter / (a[2] * a[3] + b[2] * b[3] - inter)


def offset(a, b):
    return ((a[0] + a[2] / 2 - b[0] - b[2] / 2) ** 2 + (a[1] + a[3] / 2 - b[1] - b[3] / 2) ** 2) ** 0.5


def cmd_boxes(args):
    sc = json.loads(Path(args.scenario).read_text(encoding="utf-8"))
    truth = json.loads(Path(args.truth).read_text(encoding="utf-8"))
    ok = Checks()
    ok("IoU・被覆率は 0 より大きく 1 以下、許容ずれは 0 以上",
       0 < args.min_iou <= 1 and 0 < args.min_coverage <= 1
       and math.isfinite(args.max_offset) and args.max_offset >= 0)
    images = set(args.images)
    ok("検証対象の画像名が正解データに存在する", bool(images) and images <= truth.keys(),
       ", ".join(sorted(images - truth.keys())))
    required = {tuple(pair) for pair in args.require_box}
    invalid = [(name, key) for name, key in required
               if name not in images or key not in truth.get(name, {}).get("boxes", {})]
    ok("必須の赤枠IDが検証対象画像の正解データに存在する", not invalid, str(invalid) if invalid else "")
    shots = sc.get("shots", [])
    ok("検証する shot がある", isinstance(shots, list) and bool(shots))
    if ok.failed:
        return 1
    expected = {(name, key) for name in images for key in truth[name]["boxes"] if key.startswith("data:")}
    seen = set()
    matched = set()
    for shot in shots:
        name = Path(shot.get("image", "")).name
        if name not in images:
            continue
        t = truth[name]
        seen.add(name)
        ok(f"{name}: 表示スケール {t['scale']} を指定している", abs(shot.get("scale", 1) - t["scale"]) < 0.01,
           f"scale={shot.get('scale', 1)}")
        ops = {k: v for k, v in t["boxes"].items() if k.startswith("op:")}
        for st in shot.get("steps", []):
            if st.get("kind") == "mark":
                continue
            if st.get("kind") == "data":
                if not re.fullmatch(r"[1-9][0-9]*", str(st.get("item", ""))):
                    ok(f"{name}: 取得項目No.が正の整数", False, repr(st.get("item")))
                    continue
                key = f"data:{st['item']}"
                ref = t["boxes"].get(key)
            if not st.get("box"):
                continue
            box = st["box"]
            valid = (isinstance(box, list) and len(box) == 4
                     and all(isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) for v in box)
                     and box[2] > 0 and box[3] > 0)
            ok(f"{name}: 赤枠座標が有限値で幅・高さが正", valid)
            if not valid:
                continue
            if st.get("kind") != "data":
                key, ref = max(((k, v) for k, v in ops.items()), key=lambda kv: iou(st["box"], kv[1]), default=(None, None))
            if ref is None:
                print(f"--  {name}: {key or st.get('text', '')} は正解データなし")
                continue
            u, d = iou(st["box"], ref), offset(st["box"], ref) / t["scale"]
            correct = u >= args.min_iou or d <= args.max_offset
            if correct:
                matched.add((name, key))
            detail = f"{name}: {key} の赤枠位置  … IoU {u:.2f} / 中心のずれ {d:.0f}px(CSS)"
            if st.get("kind") == "data":
                print(f"{'OK ' if correct else '-- '} {detail}")
            else:
                ok(detail, correct)
    for name in sorted(images):
        ok(f"{name} を使った shot がある", name in seen)
        ok(f"{name}: 正解と一致する赤枠がある", any(image == name for image, _ in matched))
        missing = sorted(key for image, key in expected - matched if image == name)
        if missing:
            print(f"--  {name}: 正解と一致する赤枠のない取得項目 {', '.join(missing)}")
    if expected:
        count = len(matched & expected)
        coverage = count / len(expected)
        ok("取得項目の赤枠の被覆率が基準以上", coverage >= args.min_coverage,
           f"{count}/{len(expected)} = {coverage:.1%}（基準 {args.min_coverage:.1%}）")
    for name, key in sorted(required):
        ok(f"{name}: 必須の赤枠 {key} が正解と一致している", (name, key) in matched)
    return 1 if ok.failed else 0


def cmd_actions(args):
    ok = Checks()
    risky = []
    for path in args.files:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
        for shot in data.get("shots", []):
            for st in shot.get("setup", []) + shot.get("steps", []):
                act, target = st.get("action"), json.dumps(st.get("target", ""), ensure_ascii=False)
                if isinstance(act, dict):
                    kind = act["type"] if "type" in act and len(act) > 1 else next(iter(act), "")
                else:
                    kind = act or ""
                if kind in ("fill", "type") and re.search(r"pass|パスワード", target + st.get("text", ""), re.I):
                    risky.append(f"{shot.get('id')}: パスワード欄に入力している {target}")
                if kind == "click" and re.search(r"同意|accept|agree", target, re.I):
                    risky.append(f"{shot.get('id')}: 同意ボタンを押している {target}")
                if kind in ("fill", "type") and re.search(r"login|ログイン|member_id|会員ID|username", target, re.I):
                    risky.append(f"{shot.get('id')}: ログインIDを入力している {target}")
                if kind == "click" and re.search(r"ロボット|captcha|robot", target, re.I):
                    risky.append(f"{shot.get('id')}: CAPTCHAを操作している {target}")
    ok("ログイン・同意・CAPTCHA の操作をしていない", not risky, " / ".join(risky))
    return 1 if ok.failed else 0


def main():
    ap = argparse.ArgumentParser(description="scraping-procedure-deck の生成物を機械的に確認する（目視QAの補助）")
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("summary", help="スライド構成・赤枠・番号・リンク・セクションを一覧する")
    s.add_argument("pptx")
    s.add_argument("--json", action="store_true")
    s = sub.add_parser("leaks", help="セレクタ等の技術情報やURLが文字として出ていないか")
    s.add_argument("pptx")
    s = sub.add_parser("insert", help="既存資料への差し込みが正しいか")
    s.add_argument("pptx")
    s.add_argument("--base", required=True, help="差し込み先として渡した資料（実行後のもの）")
    s.add_argument("--at", type=int, required=True, help="追加分が何枚目から入るべきか")
    s.add_argument("--hash", help="prepare.py が出力した hashes.json（元ファイルの無変更確認）")
    s = sub.add_parser("edits", help="手直し済み資料への追加で、手直しが保たれているか")
    s.add_argument("pptx")
    s.add_argument("--record", required=True, help="prepare.py が出力した hand_edits.json")
    s = sub.add_parser("contrast", help="文字と背景のコントラスト比（WCAG）")
    s.add_argument("pptx")
    s.add_argument("--slides", help="対象スライドの範囲（例: 3-9）")
    s = sub.add_parser("boxes", help="手持ち画像（C）の赤枠座標を正解と比べる")
    s.add_argument("scenario")
    s.add_argument("--truth", required=True)
    s.add_argument("--images", nargs="+", required=True, help="必須の画像名（truth.json のキー）")
    s.add_argument("--min-coverage", type=float, default=1.0, help="取得項目の正しい赤枠の割合の下限")
    s.add_argument("--require-box", nargs=2, action="append", default=[], metavar=("IMAGE", "KEY"),
                   help="割合によらず必須の赤枠（例: 会員ログイン後_商品詳細.png data:13）")
    s.add_argument("--min-iou", type=float, default=0.5)
    s.add_argument("--max-offset", type=float, default=12, help="中心のずれの許容（CSS px）")
    s = sub.add_parser("actions", help="scenario/deck.json にログイン・同意・CAPTCHA の操作がないか")
    s.add_argument("files", nargs="+")
    args = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit({"summary": cmd_summary, "leaks": cmd_leaks, "insert": cmd_insert, "edits": cmd_edits,
              "contrast": cmd_contrast, "boxes": cmd_boxes, "actions": cmd_actions}[args.cmd](args))


if __name__ == "__main__":
    main()
