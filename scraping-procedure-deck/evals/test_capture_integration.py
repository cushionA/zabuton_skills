import copy
import hashlib
import importlib.metadata
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import threading
import unittest
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from time import monotonic

from PIL import Image
from pptx import Presentation

SKILL = Path(__file__).resolve().parents[1]
try:
    importlib.metadata.version("playwright")
    HAS_PLAYWRIGHT = True
except importlib.metadata.PackageNotFoundError:
    HAS_PLAYWRIGHT = False


class QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass


@unittest.skipUnless(HAS_PLAYWRIGHT, "Playwright と Chromium が必要です")
class CaptureIntegrationTests(unittest.TestCase):
    def run_script(self, name, *args):
        return subprocess.run(
            [sys.executable, "-B", str(SKILL / "scripts" / name), *map(str, args)],
            capture_output=True, text=True, encoding="utf-8", timeout=180,
            env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        )

    def serve(self, site):
        server = ThreadingHTTPServer(("127.0.0.1", 0), partial(QuietHandler, directory=str(site)))
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        def cleanup():
            server.shutdown()
            server.server_close()
            thread.join()
        self.addCleanup(cleanup)
        return server.server_port

    def assert_green_box(self, image, box):
        image = image.convert("RGB")
        x, y, w, h = box
        green = (26, 199, 90)
        # 背景色がある四隅と、その外側を確かめ、寸法・位置の両方を検証する。
        for px, py in [(x + 3, y + 3), (x + w - 3, y + 3),
                       (x + 3, y + h - 3), (x + w - 3, y + h - 3)]:
            self.assertEqual(image.getpixel((round(px), round(py))), green)
        for px, py in [(x - 3, y + h - 5), (x + w + 3, y + h - 5),
                       (x + w - 5, y - 3), (x + w - 5, y + h + 3)]:
            self.assertNotEqual(image.getpixel((round(px), round(py))), green)

    def test_delayed_values_and_fractional_scale_match_saved_pixels(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            (work / "late.html").write_text("""<meta charset="utf-8"><style>
              body {margin:0; height:1000px} #value {position:absolute;left:51px;top:121px;
              width:121px;height:41px;background:rgb(26,199,90)}</style>
              <body><script>setTimeout(() => {
                document.body.insertAdjacentHTML('beforeend', '<div id="value"></div>');
                setTimeout(() => { const v = document.querySelector('#value');
                  v.style.top='151px'; v.textContent='準備完了'; }, 600);
              }, 800);</script>""", encoding="utf-8")
            url = f"http://127.0.0.1:{self.serve(work)}/late.html"
            scenario = {"title": "遅延表示", "data_layout": [{"no": 1, "name": "値", "type": "文字列"}],
                        "browser": {"viewport": [901, 601], "device_scale_factor": 1.5,
                                    "settle_ms": 0, "timeout_ms": 6000},
                        "shots": [{"id": "late", "screen": "値", "goto": url,
                                   "wait_for": {"css": "#value", "has_text": "準備完了"},
                                   "steps": [{"kind": "data", "item": 1, "target": "#value", "fit": "box"}],
                                   "peeks": [{"goto": url, "target": "#value", "mark": "#value", "pad": 8.5,
                                              "text": "表示差分", "timeout_ms": 6000,
                                              "wait_for": {"css": "#value", "has_text": "準備完了"}}]}]}
            source = work / "scenario.json"
            source.write_text(json.dumps(scenario, ensure_ascii=False), encoding="utf-8")
            for attempt in range(3):
                with self.subTest(attempt=attempt):
                    captures = work / f"captures{attempt}"
                    result = self.run_script("capture.py", source, "--out", captures)
                    self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
                    shot = json.loads((captures / "deck.json").read_text(encoding="utf-8"))["shots"][0]
                    self.assertEqual(shot["steps"][0]["box"], [76.5, 226.5, 181.5, 61.5])
                    for info, box in [(shot, shot["steps"][0]["box"]),
                                      (shot["peeks"][0], shot["peeks"][0]["mark_box"])]:
                        with Image.open(captures / info["image"]) as image:
                            self.assertEqual(info["image_size"], list(image.size))
                            self.assert_green_box(image, box)
                    self.assertEqual(shot["image_size"][0], 1352)
            output = work / "new" / "nested" / "deck.pptx"
            built = self.run_script("build_deck.py", source, "--captures", captures, "-o", output)
            self.assertEqual(built.returncode, 0, built.stdout + built.stderr)
            self.assertGreaterEqual(len(Presentation(output).slides), 4)

    def test_nested_cross_origin_frames_text_and_image_map_coordinates(self):
        from playwright.sync_api import sync_playwright
        spec = importlib.util.spec_from_file_location("capture_real", SKILL / "scripts" / "capture.py")
        capture = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(capture)
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            port = self.serve(work)
            (work / "index.html").write_text(f"""<style>body {{margin:0;height:2000px}}
              iframe {{position:absolute;left:31px;top:151px;border:7px solid black}}
              </style><iframe id="outer" width="700" height="500"
              src="http://localhost:{port}/outer.html"></iframe>""", encoding="utf-8")
            (work / "outer.html").write_text(f"""<style>body {{margin:0;height:1500px}}
              iframe {{position:absolute;left:17px;top:90px;border:5px solid black}}
              </style><iframe id="inner" width="600" height="350"
              src="http://127.0.0.1:{port}/inner.html"></iframe>""", encoding="utf-8")
            (work / "inner.html").write_text("""<meta charset="utf-8"><style>body {margin:0;height:1200px}
              #panel {position:absolute;left:13.25px;top:97.5px;width:301px;height:150px;background:#ddd}
              #cell {position:absolute;left:17.25px;top:19.5px;width:91.5px;height:31px;background:rgb(26,199,90)}
              #text {position:absolute;left:7px;top:80px;padding:8px;font:16px sans-serif}
              img {position:absolute;left:27px;top:300px;border:3px solid black}
              </style><section id="panel"><div id="cell"></div><span id="text">在庫あり</span></section>
              <img width="200" height="100" usemap="#stock"
                src="data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' width='200' height='100'%3E%3Crect width='200' height='100' fill='orange'/%3E%3C/svg%3E">
              <map name="stock"><area id="area" shape="rect" coords="30,20,80,60" href="#">
              <area id="circle" shape="circle" coords="60,40,10" href="#">
              <area id="poly" shape="poly" coords="10,10,40,10,20,30" href="#">
              <area id="whole" shape="default" href="#"></map>""", encoding="utf-8")
            with sync_playwright() as p:
                browser = p.chromium.launch()
                ctx = browser.new_context(viewport={"width": 901, "height": 601}, device_scale_factor=1.5)
                page = ctx.new_page()
                page.goto(f"http://127.0.0.1:{port}/index.html")
                frames = ["#outer", "#inner"]
                cell_target = {"frame": frames, "css": "#cell", "within": {"css": "#panel"}}
                self.assertEqual(capture.measure(page, cell_target, 5000, "box"), [90.5, 370, 91.5, 31])
                outer = page.frames[1]
                inner = page.frames[2]
                page.evaluate("window.scrollTo(0, 100)")
                outer.evaluate("window.scrollTo(0, 30)")
                inner.evaluate("window.scrollTo(0, 40)")
                self.assertEqual(capture.measure(page, cell_target, 5000, "box"), [90.5, 300, 91.5, 31])
                text_target = {"frame": frames, "css": "#text"}
                text_box = capture.measure(page, text_target, 5000, "box")
                text_rect = capture.measure(page, text_target, 5000, "text")
                self.assertAlmostEqual(text_rect[0] - text_box[0], 8)
                self.assertLess(text_rect[2], text_box[2])
                self.assertLess(text_rect[3], text_box[3])
                for selector, expected in [("#area", [120, 506, 50, 40]),
                                           ("#circle", [140, 516, 20, 20]),
                                           ("#poly", [100, 496, 30, 20]),
                                           ("#whole", [90, 486, 200, 100])]:
                    with self.subTest(selector=selector):
                        self.assertEqual(capture.measure(page, {"frame": frames, "css": selector}, 5000, "box"), expected)
                page.evaluate("window.scrollTo(0, 0)")
                outer.evaluate("window.scrollTo(0, 0)")
                inner.evaluate("window.scrollTo(0, 0)")
                png = work / "frame.png"
                page.screenshot(path=str(png))
                with Image.open(png) as image:
                    self.assertEqual(image.convert("RGB").getpixel((round(100 * 1.5), round(380 * 1.5))),
                                     (26, 199, 90))
                peek = {"goto": page.url, "target": {"frame": frames, "css": "#panel"},
                        "mark": cell_target, "pad": 8.5}
                capture.capture_peek(ctx, peek, work / "peek.png", 1.5, 601, 5000, 0)
                with Image.open(work / "peek.png") as image:
                    self.assertEqual(peek["image_size"], list(image.size))
                    self.assert_green_box(image, peek["mark_box"])
                started = monotonic()
                with self.assertRaisesRegex(capture.PlaywrightError, "750ms"):
                    capture.stable(page, lambda ms: [capture.measure(page, "#missing", ms, "box")], 750)
                self.assertLess(monotonic() - started, 1.5)
                browser.close()

    def test_partial_recapture_and_build_with_real_browser(self):
        site = SKILL / "evals" / "files" / "mock-shop" / "site"
        server = ThreadingHTTPServer(("127.0.0.1", 0), partial(QuietHandler, directory=str(site)))
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            with tempfile.TemporaryDirectory() as tmp:
                work = Path(tmp)
                source = SKILL / "evals" / "files" / "mock-shop" / "scenario.json"
                base_url = f"http://127.0.0.1:{server.server_port}/"
                scenario = json.loads(source.read_text(encoding="utf-8").replace("http://localhost:8765/", base_url))
                scenario["shots"][0]["steps"][0]["action"] = {"type": "ワイヤレスイヤホン"}
                scenario_path = work / "scenario.json"
                scenario_path.write_text(json.dumps(scenario, ensure_ascii=False), encoding="utf-8")
                captures = work / "captures"
                first = self.run_script("capture.py", scenario_path, "--out", captures)
                self.assertEqual(first.returncode, 0, first.stdout + first.stderr)
                original = json.loads((captures / "deck.json").read_text(encoding="utf-8"))
                original_images = {
                    shot["id"]: (
                        hashlib.sha256((captures / shot["image"]).read_bytes()).hexdigest(),
                        (captures / shot["image"]).stat().st_mtime_ns,
                    )
                    for shot in original["shots"][:3]
                }
                price_step = next(step for step in scenario["shots"][-1]["steps"] if step.get("item") == 4)
                price_step["target"] = "[class*='Price_price']"
                scenario_path.write_text(json.dumps(scenario, ensure_ascii=False), encoding="utf-8")
                result = self.run_script("capture.py", scenario_path, "--out", captures, "--only", "item")
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
                updated = json.loads((captures / "deck.json").read_text(encoding="utf-8"))
                self.assertEqual(original["shots"][:3], updated["shots"][:3])
                for shot in updated["shots"][:3]:
                    image = captures / shot["image"]
                    self.assertEqual(original_images[shot["id"]], (
                        hashlib.sha256(image.read_bytes()).hexdigest(), image.stat().st_mtime_ns,
                    ))
                old_item, new_item = original["shots"][-1], updated["shots"][-1]
                self.assertNotEqual(old_item["captured_at"], new_item["captured_at"])
                old_price = next(step for step in old_item["steps"] if step.get("item") == 4)
                new_price = next(step for step in new_item["steps"] if step.get("item") == 4)
                self.assertEqual(new_price["target"], price_step["target"])
                self.assertNotEqual(old_price["box"], new_price["box"])
                self.assertIn("item.html?id=1001", new_item["captured_url"])
                output = work / "deck.pptx"
                built = self.run_script("build_deck.py", scenario_path, "--captures", captures, "-o", output)
                self.assertEqual(built.returncode, 0, built.stdout + built.stderr)
                self.assertGreaterEqual(len(Presentation(output).slides), 5)
                reordered = copy.deepcopy(scenario)
                reordered["shots"][-1]["peeks"].reverse()
                scenario_path.write_text(json.dumps(reordered, ensure_ascii=False), encoding="utf-8")
                rejected_output = work / "mismatched.pptx"
                rejected = self.run_script("build_deck.py", scenario_path, "--captures", captures, "-o", rejected_output)
                self.assertNotEqual(rejected.returncode, 0)
                self.assertFalse(rejected_output.exists())
        finally:
            server.shutdown()
            server.server_close()
            thread.join()


if __name__ == "__main__":
    unittest.main()
