import argparse
import base64
import copy
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path, PurePosixPath
from unittest.mock import patch

from PIL import Image, ImageDraw
from pptx import Presentation


SCRIPTS = Path(__file__).parents[1] / "scripts"


def load_module(name):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


importer = load_module("import_results")
builder_module = load_module("build_deck")


class ImportResultsTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.artifacts = self.root / "artifacts"
        self.artifacts.mkdir()
        self.report_path = self.artifacts / "report.json"
        self.plan_path = self.root / "plan.json"
        self.out = self.root / "imported"
        image = Image.new("RGB", (1280, 800), "white")
        ImageDraw.Draw(image).rectangle((80, 80, 280, 120), fill="navy")
        image.save(self.artifacts / "selected.png")
        self.png = (self.artifacts / "selected.png").read_bytes()
        self.report = {
            "config": {"projects": [{"id": "chromium", "name": "chromium",
                                       "outputDir": str(self.artifacts)}], "secret": "CONFIG_SECRET"},
            "suites": [{"title": "flow.spec.ts", "suites": [{"title": "商品確認", "specs": [{
                "title": "価格", "ok": True, "tests": [{"projectName": "chromium", "projectId": "chromium",
                    "results": [{"status": "passed", "retry": 0,
                                 "stdout": [{"text": "STDOUT_SECRET"}], "errors": [{"message": "ERROR_SECRET"}],
                                 "attachments": [{"name": "価格", "contentType": "image/png", "path": "selected.png"},
                                                 {"name": "trace", "contentType": "application/zip", "path": "trace.zip"}]}]}]}]}]}],
        }
        self.plan = {"title": "商品確認手順", "site": "テスト店舗", "data_layout": [
            {"no": 1, "name": "価格", "type": "数値"},
            {"no": 2, "name": "処理日時", "type": "日時", "on_screen": False, "origin": "処理時に付与"},
        ], "shots": [{"attachment": "r001", "id": "price", "screen": "商品詳細",
                      "steps": [{"kind": "data", "item": 1, "box": [80, 80, 200, 40]}]}]}

    def result(self):
        return self.report["suites"][0]["suites"][0]["specs"][0]["tests"][0]["results"][0]

    def record(self):
        return self.report["suites"][0]["suites"][0]["specs"][0]["tests"][0]

    def save_inputs(self):
        self.report_path.write_text(json.dumps(self.report, ensure_ascii=False), encoding="utf-8")
        self.plan_path.write_text(json.dumps(self.plan, ensure_ascii=False), encoding="utf-8")

    def run_import(self, artifacts=None):
        self.save_inputs()
        return importer.import_results(self.report_path, self.plan_path, self.out, artifacts)

    def assert_rejected(self, pattern=None, artifacts=None):
        with self.assertRaises(importer.ImportError) as raised:
            self.run_import(artifacts)
        if pattern:
            self.assertIn(pattern, str(raised.exception))
        self.assertFalse(self.out.exists())

    def test_list_disambiguates_nested_tests_projects_retries_without_reading_images(self):
        test = self.record()
        test["results"].append(copy.deepcopy(test["results"][0]))
        test["results"][0].update(status="failed")
        test["results"][1].update(retry=1)
        second = copy.deepcopy(test)
        second.update(projectName="firefox", projectId="firefox")
        self.report["suites"][0]["suites"][0]["specs"][0]["tests"].append(second)
        with patch.object(importer, "read_png", side_effect=AssertionError("画像を開いた")):
            candidates = importer.list_candidates(self.report)["attachments"]
        self.assertEqual([c["id"] for c in candidates], ["r001", "r002", "r003", "r004"])
        self.assertEqual([c["project"] for c in candidates], ["chromium", "chromium", "firefox", "firefox"])
        self.assertEqual([c["retry"] for c in candidates], [0, 1, 0, 1])
        self.assertEqual(candidates[0]["test"], "flow.spec.ts / 商品確認 / 価格")
        serialized = json.dumps(candidates)
        for secret in ("CONFIG_SECRET", "STDOUT_SECRET", "ERROR_SECRET", "trace.zip", "_attachment"):
            self.assertNotIn(secret, serialized)

    def test_import_uses_only_selected_png_and_plan_order(self):
        self.result()["attachments"].extend([
            {"name": "未選択", "contentType": "image/png", "path": "missing.png"},
            {"name": "別画像", "contentType": "image/png", "body": base64.b64encode(self.png).decode()},
        ])
        second = copy.deepcopy(self.plan["shots"][0])
        second.update(attachment="r003", id="second")
        self.plan["shots"].insert(0, second)
        with patch.object(importer, "read_png", wraps=importer.read_png) as reader:
            self.assertEqual(self.run_import(), [])
        self.assertEqual([c.args[1]["id"] for c in reader.call_args_list], ["r003", "r001"])
        deck = json.loads((self.out / "scenario.json").read_text(encoding="utf-8"))
        self.assertEqual([s["id"] for s in deck["shots"]], ["second", "price"])
        self.assertEqual(deck["shots"][0]["image_size"], [1280, 800])
        self.assertEqual({p.name for p in (self.out / "images").iterdir()}, {"shot_second.png", "shot_price.png"})
        self.assertEqual((self.out / "images" / "shot_price.png").read_bytes(), self.png)

    def test_retry_success_warns_and_does_not_use_failed_attempt(self):
        failed = copy.deepcopy(self.result())
        failed.update(status="failed")
        failed["attachments"][0]["path"] = "missing-failure.png"
        self.result()["retry"] = 1
        self.record()["results"].insert(0, failed)
        self.plan["shots"][0]["attachment"] = "r002"
        self.assertEqual(self.run_import(), ["r002: retry=1 の成功結果を使用します"])

    def test_non_passed_result_rejected_even_if_spec_is_ok(self):
        for status in ("failed", "timedOut", "skipped", "interrupted"):
            with self.subTest(status=status):
                self.result()["status"] = status
                self.assert_rejected("passed")

    def test_inline_body_is_supported_and_not_listed(self):
        self.result()["attachments"][0] = {"name": "inline", "contentType": "image/png",
                                            "body": base64.b64encode(self.png).decode()}
        candidate = importer.list_candidates(self.report)["attachments"][0]
        self.assertIsNone(candidate["path"])
        self.assertNotIn("body", candidate)
        self.run_import()
        self.assertEqual((self.out / "images" / "shot_price.png").read_bytes(), self.png)

    def test_invalid_base64_body_has_clear_error(self):
        for body in ("%%%", "日本語"):
            with self.subTest(body=body):
                self.result()["attachments"][0] = {"name": "bad", "contentType": "image/png", "body": body}
                self.assert_rejected("base64")

    def test_all_selected_images_validated_before_creating_output(self):
        self.result()["attachments"].append({"name": "bad", "contentType": "image/png", "path": "absent.png"})
        second = copy.deepcopy(self.plan["shots"][0])
        second.update(attachment="r002", id="last")
        self.plan["shots"].append(second)
        self.assert_rejected("読み込めません")

    def test_existing_output_is_never_overwritten(self):
        self.out.mkdir()
        marker = self.out / "keep.txt"
        marker.write_text("keep")
        with self.assertRaisesRegex(importer.ImportError, "新規フォルダ"):
            self.run_import()
        self.assertEqual(marker.read_text(), "keep")
        self.assertEqual(list(self.out.iterdir()), [marker])

    def test_empty_or_unknown_selection_is_rejected(self):
        self.plan["shots"] = []
        self.assert_rejected("1 件以上")
        self.plan["shots"] = [{"attachment": "r999"}]
        self.assert_rejected("候補がありません")

    def test_duplicate_selection_and_ids_are_rejected(self):
        self.plan["shots"].append(copy.deepcopy(self.plan["shots"][0]))
        self.assert_rejected("attachment が重複")
        self.result()["attachments"].append(copy.deepcopy(self.result()["attachments"][0]))
        self.plan["shots"][1].update(attachment="r002", id="PRICE")
        self.assert_rejected("id が重複")

    def test_unsafe_ids_are_rejected_and_reserved_names_are_prefixed(self):
        for sid in ("../out", "/abs", "日本語", "name.", "name:stream", "_start", "a" * 81):
            with self.subTest(sid=sid):
                self.plan["shots"][0]["id"] = sid
                self.assert_rejected("ASCII")
        self.plan["shots"][0]["id"] = "CON"
        self.run_import()
        self.assertTrue((self.out / "images" / "shot_CON.png").exists())

    def test_path_traversal_and_external_absolute_paths_are_rejected(self):
        outside = self.root / "outside.png"
        outside.write_bytes(self.png)
        for path in ("../outside.png", "..\\outside.png", str(outside), "C:outside.png", "\\outside.png"):
            with self.subTest(path=path):
                self.result()["attachments"][0]["path"] = path
                self.assert_rejected()

    def test_symlink_cannot_escape_artifacts(self):
        outside = self.root / "outside.png"
        outside.write_bytes(self.png)
        link = self.artifacts / "link.png"
        try:
            os.symlink(outside, link)
        except OSError as exc:
            self.skipTest(f"シンボリックリンクを作成できません: {exc}")
        self.result()["attachments"][0]["path"] = "link.png"
        self.assert_rejected("外")

    def test_absolute_path_inside_artifacts_is_allowed(self):
        self.result()["attachments"][0]["path"] = str(self.artifacts / "selected.png")
        self.run_import()
        self.assertTrue((self.out / "scenario.json").exists())

    def test_relocated_windows_report_maps_from_project_output_dir(self):
        self.report["config"]["projects"][0]["outputDir"] = "C:/old/project/test-results"
        self.result()["attachments"][0]["path"] = "c:\\old\\project\\test-results\\case\\attachments\\shot.png"
        relocated = self.artifacts / "case" / "attachments"
        relocated.mkdir(parents=True)
        (relocated / "shot.png").write_bytes(self.png)
        self.run_import(self.artifacts)
        self.assertEqual((self.out / "images" / "shot_price.png").read_bytes(), self.png)

    def test_relocated_posix_report_maps_from_project_output_dir(self):
        self.report["config"]["projects"][0]["outputDir"] = "/old/project/test-results"
        self.result()["attachments"][0]["path"] = "/old/project/test-results/selected.png"
        self.run_import(self.artifacts)
        self.assertTrue((self.out / "scenario.json").exists())

    def test_foreign_absolute_path_is_not_treated_as_native_relative(self):
        self.report["config"]["projects"][0]["outputDir"] = "C:/old/test-results"
        self.result()["attachments"][0]["path"] = "C:\\old\\test-results\\selected.png"
        candidate = importer.collect_candidates(self.report)[0]
        with patch.object(importer, "Path", PurePosixPath):
            selected = importer.select_candidates(self.report, ["r001"], self.artifacts, True)
            mapped = importer.attachment_path(self.report, candidate, self.artifacts, True)
        self.assertEqual(selected[0]["id"], "r001")
        self.assertEqual(mapped, (self.artifacts / "selected.png").resolve())

    def test_remap_requires_original_path_under_project_output_dir(self):
        self.report["config"]["projects"][0]["outputDir"] = "C:/old/test-results"
        self.result()["attachments"][0]["path"] = "C:/elsewhere/selected.png"
        self.assert_rejected("outputDir の外", self.artifacts)

    def test_project_id_disambiguates_duplicate_names_for_remapping(self):
        self.report["config"]["projects"] = [
            {"id": "chromium", "name": "same", "outputDir": "/chosen/results"},
            {"id": "other", "name": "same", "outputDir": "/wrong/results"},
        ]
        self.record()["projectName"] = "same"
        self.result()["attachments"][0]["path"] = "/chosen/results/selected.png"
        self.run_import(self.artifacts)
        self.assertTrue((self.out / "scenario.json").exists())

    def test_ambiguous_project_without_id_is_rejected(self):
        self.record().pop("projectId")
        self.report["config"]["projects"] *= 2
        self.result()["attachments"][0]["path"] = "/old/selected.png"
        self.assert_rejected("一意", self.artifacts)

    def test_different_project_output_dirs_cannot_share_one_relocation_root(self):
        self.report["config"]["projects"][0]["outputDir"] = "/old/chromium-results"
        self.result()["attachments"][0]["path"] = "/old/chromium-results/selected.png"
        second_test = copy.deepcopy(self.record())
        second_test.update(projectId="firefox", projectName="firefox")
        second_test["results"][0]["attachments"][0]["path"] = "/old/firefox-results/selected.png"
        self.report["suites"][0]["suites"][0]["specs"][0]["tests"].append(second_test)
        self.report["config"]["projects"].append({"id": "firefox", "name": "firefox", "outputDir": "/old/firefox-results"})
        second_shot = copy.deepcopy(self.plan["shots"][0])
        second_shot.update(id="second", attachment="r002")
        self.plan["shots"].append(second_shot)
        self.assert_rejected("分けて", self.artifacts)

    def test_corrupt_image_and_wrong_format_are_rejected(self):
        (self.artifacts / "selected.png").write_bytes(b"not a PNG")
        self.assert_rejected("デコード")
        Image.new("RGB", (100, 100)).save(self.artifacts / "selected.png", format="JPEG")
        self.assert_rejected("PNG ではありません")

    def test_rectangles_are_finite_positive_and_inside_the_actual_image(self):
        invalid = [[0, 0, 0, 1], [0, 0, -1, 2], [-1, 0, 2, 2], [0, -1, 2, 2],
                   [1279, 0, 2, 1], [0, 799, 1, 2], [0, 0, float("nan"), 1],
                   [0, 0, float("inf"), 1], [True, 0, 1, 1], [0, 0, 1], None]
        for box in invalid:
            with self.subTest(box=box):
                self.plan["shots"][0]["steps"][0]["box"] = box
                self.assert_rejected("box")

    def test_scale_must_be_positive_and_finite(self):
        for scale in (0, -1, float("nan"), float("inf"), True, "2", 10 ** 400):
            with self.subTest(scale=scale):
                self.plan["shots"][0]["scale"] = scale
                self.assert_rejected("scale")

    def test_data_layout_numbers_and_references_are_validated(self):
        original = copy.deepcopy(self.plan)
        for no in (0, -1, True, 1.5):
            with self.subTest(no=no):
                self.plan = copy.deepcopy(original)
                self.plan["data_layout"][0]["no"] = no
                self.assert_rejected("no")
        self.plan = copy.deepcopy(original)
        self.plan["data_layout"].append(copy.deepcopy(self.plan["data_layout"][0]))
        self.assert_rejected("重複")
        self.plan = copy.deepcopy(original)
        self.plan["shots"][0]["steps"][0]["item"] = 9
        self.assert_rejected("対応する")
        self.plan["shots"][0]["steps"][0]["item"] = 2
        self.assert_rejected("画面外")

    def test_step_kinds_text_and_boxes_are_required(self):
        invalid = [[], [{"kind": "click", "box": [0, 0, 1, 1]}],
                   [{"kind": "op", "box": [0, 0, 1, 1]}],
                   [{"kind": "mark", "text": "", "box": [0, 0, 1, 1]}],
                   [{"kind": "op", "text": "進む"}], [{"kind": "mark", "text": "対象外"}]]
        for steps in invalid:
            with self.subTest(steps=steps):
                self.plan["shots"][0]["steps"] = steps
                self.assert_rejected()

    def test_source_execution_and_link_fields_are_not_copied(self):
        self.plan.update(start_url="https://secret.invalid", browser={"headers": "SECRET"})
        shot = self.plan["shots"][0]
        shot.update(url="https://secret.invalid", goto="https://secret.invalid", link=True,
                    captured_url="https://secret.invalid", setup=[{"action": "click"}])
        shot["steps"][0].update(target="#secret", action="click", text="IGNORED_SOURCE_TEXT")
        self.run_import()
        raw = (self.out / "scenario.json").read_text(encoding="utf-8")
        self.assertFalse(json.loads(raw)["shots"][0]["link"])
        for value in ("secret.invalid", "SECRET", "#secret", "click", "IGNORED_SOURCE_TEXT"):
            self.assertNotIn(value, raw)

    def test_type_is_required_and_notes_are_text_arrays(self):
        self.plan["data_layout"][0].pop("type")
        self.assert_rejected("type")
        self.plan["data_layout"][0]["type"] = "数値"
        self.plan["shots"][0]["notes"] = "配列ではない"
        self.assert_rejected("notes")
        self.plan["shots"][0]["notes"] = ["次の商品の確認を繰り返す"]
        self.run_import()
        deck = json.loads((self.out / "scenario.json").read_text(encoding="utf-8"))
        self.assertEqual(deck["shots"][0]["notes"], ["次の商品の確認を繰り返す"])

    def test_extract_only_selected_inline_image_with_retry_and_minimal_index(self):
        self.result()["retry"] = 1
        self.result()["attachments"] = [
            {"name": "未選択", "contentType": "image/png", "path": "missing.png"},
            {"name": "選択", "contentType": "image/png", "body": base64.b64encode(self.png).decode()},
        ]
        self.save_inputs()
        with patch.object(importer, "read_png", wraps=importer.read_png) as reader:
            warnings = importer.extract_results(self.report_path, ["r002"], self.out)
        self.assertEqual([call.args[1]["id"] for call in reader.call_args_list], ["r002"])
        self.assertEqual(warnings, ["r002: retry=1 の成功結果を使用します"])
        self.assertEqual({p.name for p in self.out.iterdir()}, {"r002.png", "index.json"})
        self.assertEqual((self.out / "r002.png").read_bytes(), self.png)
        index = json.loads((self.out / "index.json").read_text(encoding="utf-8"))
        self.assertEqual(index["images"][0]["image_size"], [1280, 800])
        self.assertEqual(index["images"][0]["id"], "r002")
        self.assertNotIn("body", index["images"][0])

    def test_extract_rejects_unknown_duplicate_failed_and_missing_selections_atomically(self):
        self.result()["attachments"].append({"name": "欠落", "contentType": "image/png", "path": "missing.png"})
        self.save_inputs()
        for selection in ([], ["r999"], ["r001", "r001"], ["r001", "r002"]):
            with self.subTest(selection=selection):
                with self.assertRaises(importer.ImportError):
                    importer.extract_results(self.report_path, selection, self.out)
                self.assertFalse(self.out.exists())
        self.result()["status"] = "failed"
        self.save_inputs()
        with self.assertRaisesRegex(importer.ImportError, "passed"):
            importer.extract_results(self.report_path, ["r001"], self.out)
        self.assertFalse(self.out.exists())

    def test_extract_cli_supports_relocated_png_and_refuses_overwrite(self):
        self.report["config"]["projects"][0]["outputDir"] = "/old/results"
        self.result()["attachments"][0]["path"] = "/old/results/selected.png"
        self.save_inputs()
        command = [sys.executable, "-X", "utf8", "-B", str(SCRIPTS / "import_results.py"), str(self.report_path),
                   "--extract", "r001", "--out", str(self.out), "--artifacts", str(self.artifacts)]
        result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", check=False)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual((self.out / "r001.png").read_bytes(), self.png)
        again = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", check=False)
        self.assertEqual(again.returncode, 2)
        self.assertNotIn("Traceback", again.stderr)
        self.assertEqual((self.out / "r001.png").read_bytes(), self.png)

    def test_imported_scenario_builds_editable_pptx(self):
        self.plan["shots"][0]["steps"].extend([
            {"kind": "op", "text": "次の画面へ進む", "box": [80, 170, 200, 40]},
            {"kind": "mark", "text": "参考価格は取得しない", "box": [80, 250, 200, 40]},
        ])
        self.run_import()
        deck = json.loads((self.out / "scenario.json").read_text(encoding="utf-8"))
        args = argparse.Namespace(base=None, only=None, insert_at=None,
                                  sections="cover,overview,steps,summary", max_per_slide=5)
        builder = builder_module.Builder(deck, self.out, args)
        destination = self.out / "deck.pptx"
        builder.build(destination)
        self.assertEqual(builder.overflows, [])
        self.assertEqual(builder.warnings, [])
        presentation = Presentation(destination)
        self.assertEqual(len(presentation.slides), 5)
        shapes = [shape for slide in presentation.slides for shape in slide.shapes]
        self.assertTrue(any(shape.has_text_frame and "次の画面へ進む" in shape.text for shape in shapes))
        self.assertTrue(any(shape.has_text_frame and "参考価格は取得しない" in shape.text for shape in shapes))
        self.assertTrue(any(shape.shape_type == 13 for shape in shapes))
        self.assertTrue(any(shape.shape_type == 1 for shape in shapes))

    def test_cli_lists_and_imports_without_browser(self):
        self.save_inputs()
        command = [sys.executable, "-X", "utf8", "-B", str(SCRIPTS / "import_results.py"), str(self.report_path)]
        listed = subprocess.run(command + ["--list"], capture_output=True, text=True, encoding="utf-8", check=False)
        self.assertEqual(listed.returncode, 0, listed.stderr)
        self.assertEqual(json.loads(listed.stdout)["attachments"][0]["id"], "r001")
        imported = subprocess.run(command + ["--plan", str(self.plan_path), "--out", str(self.out)],
                                  capture_output=True, text=True, encoding="utf-8", check=False)
        self.assertEqual(imported.returncode, 0, imported.stderr)
        self.assertTrue((self.out / "scenario.json").exists())
        failed = subprocess.run(command + ["--plan", str(self.plan_path), "--out", str(self.out)],
                                capture_output=True, text=True, encoding="utf-8", check=False)
        self.assertEqual(failed.returncode, 2)
        self.assertNotIn("Traceback", failed.stderr)


if __name__ == "__main__":
    unittest.main()

