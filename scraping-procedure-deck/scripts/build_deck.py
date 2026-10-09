import argparse
import copy
import json
import os
import re
import sys
import uuid
from datetime import date
from pathlib import Path

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

DEFAULT_FONT = "Meiryo UI"
FONT_FILES = {
    "Meiryo UI": (("meiryo.ttc", 2), ("meiryob.ttc", 2)),
    "Meiryo": (("meiryo.ttc", 0), ("meiryob.ttc", 0)),
    "メイリオ": (("meiryo.ttc", 0), ("meiryob.ttc", 0)),
    "Yu Gothic": (("YuGothM.ttc", 0), ("YuGothB.ttc", 0)),
    "游ゴシック": (("YuGothM.ttc", 0), ("YuGothB.ttc", 0)),
    "Yu Gothic UI": (("YuGothM.ttc", 1), ("YuGothB.ttc", 1)),
    "BIZ UDPGothic": (("BIZ-UDGothicR.ttc", 1), ("BIZ-UDGothicB.ttc", 1)),
    "BIZ UDPゴシック": (("BIZ-UDGothicR.ttc", 1), ("BIZ-UDGothicB.ttc", 1)),
    "BIZ UDGothic": (("BIZ-UDGothicR.ttc", 0), ("BIZ-UDGothicB.ttc", 0)),
    "BIZ UDゴシック": (("BIZ-UDGothicR.ttc", 0), ("BIZ-UDGothicB.ttc", 0)),
    "MS PGothic": (("msgothic.ttc", 2), ("msgothic.ttc", 2)),
    "ＭＳ Ｐゴシック": (("msgothic.ttc", 2), ("msgothic.ttc", 2)),
    "MS Gothic": (("msgothic.ttc", 0), ("msgothic.ttc", 0)),
    "ＭＳ ゴシック": (("msgothic.ttc", 0), ("msgothic.ttc", 0)),
    "Noto Sans JP": (("NotoSansJP-VF.ttf", 0), ("NotoSansJP-VF.ttf", 0)),
}
LINE_FACTOR = 1.3

RED = RGBColor(0xFF, 0x00, 0x00)
GRAY_BG = RGBColor(0xD9, 0xD9, 0xD9)
GRAY_TX = RGBColor(0x7F, 0x7F, 0x7F)
HIT_BG = RGBColor(0xFF, 0xEE, 0xEE)
RULE = RGBColor(0xC9, 0xCE, 0xD6)
T = MSO_THEME_COLOR
PALETTES = {
    False: {"ink": RGBColor(0x26, 0x2A, 0x30), "muted": RGBColor(0x5F, 0x66, 0x70),
            "panel": RGBColor(0xF3, 0xF4, 0xF6), "white": RGBColor(0xFF, 0xFF, 0xFF),
            "nav": RGBColor(0x37, 0x41, 0x51), "data": RGBColor(0x0E, 0x74, 0x90),
            "head": RGBColor(0x37, 0x41, 0x51), "link": RGBColor(0x05, 0x63, 0xC1)},
    # 既存資料に差し込むときはテーマ色を参照し、資料側の配色に従わせる
    True: {"ink": (T.TEXT_1, 0), "muted": (T.TEXT_1, 0.4), "panel": (T.BACKGROUND_1, -0.05),
           "white": (T.BACKGROUND_1, 0), "nav": (T.TEXT_2, 0), "data": (T.ACCENT_1, 0),
           "head": (T.TEXT_2, 0), "link": (T.HYPERLINK, 0)},
}

LABEL = 0.3
MIN_FRAME = (0.2, 0.16)
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
SECTION_EXT = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
MIN_CROP_CSS_W = 1000
GRAY_NOTE = "※グレーの項目は画面から取得しない項目です"
FOOTER_TYPES = (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE)


def circled(n):
    if 1 <= n <= 20:
        return chr(0x2460 + n - 1)
    if 21 <= n <= 35:
        return chr(0x3251 + n - 21)
    if 36 <= n <= 50:
        return chr(0x32B1 + n - 36)
    return f"({n})"


class Metrics:
    def __init__(self, name):
        self.files = FONT_FILES.get(name, FONT_FILES[DEFAULT_FONT])
        self.factor = 1.0 if name in FONT_FILES else 1.08
        self.cache = {}
        self.dirs = [Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts",
                     Path.home() / "AppData/Local/Microsoft/Windows/Fonts",
                     Path("/Library/Fonts"), Path.home() / "Library/Fonts"]

    def _font(self, pt, bold):
        key = (pt, bold)
        if key not in self.cache:
            name, index = self.files[1 if bold else 0]
            self.cache[key] = next((ImageFont.truetype(str(d / name), int(pt * 10), index=index)
                                    for d in self.dirs if (d / name).exists()), None)
        return self.cache[key]

    def width(self, text, pt, bold=False):
        f = self._font(pt, bold)
        if f is None:
            return sum(pt * (1.0 if ord(c) > 0x2E7F else 0.6) for c in text) * self.factor
        return f.getlength(text) / 10 * self.factor

    def lines(self, text, pt, width_pt, bold=False):
        return sum(self._wrap(para, pt, width_pt, bold) for para in text.split("\n"))

    def _wrap(self, text, pt, width_pt, bold):
        lines, cur = 1, 0.0
        for tok in re.findall(r"[A-Za-z0-9_\-.,:;/()%&+#@'\"!?\[\]]+|\s+|.", text):
            w = self.width(tok, pt, bold)
            if cur + w <= width_pt:
                cur += w
            elif tok.isspace():
                lines, cur = lines + 1, 0.0
            elif w > width_pt:
                for ch in tok:
                    cw = self.width(ch, pt, bold)
                    if cur + cw > width_pt and cur > 0:
                        lines, cur = lines + 1, 0.0
                    cur += cw
            else:
                lines, cur = lines + 1, w
        return lines

    def height(self, text, pt, width_in, bold=False):
        return self.lines(text, pt, width_in * 72, bold) * pt * LINE_FACTOR / 72


M = Metrics(DEFAULT_FONT)
FONT_NAME = DEFAULT_FONT
C = PALETTES[False]


def apply_color(cf, c):
    if isinstance(c, RGBColor):
        cf.rgb = c
    else:
        cf.theme_color = c[0]
        if c[1]:
            cf.brightness = c[1]


def style_run(run, size=None, color=None, bold=None, url=None):
    f = run.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        apply_color(f.color, color)
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "ja-JP")
    rpr.set("altLang", "en-US")
    if FONT_NAME:
        f.name = FONT_NAME
        prev = rpr.find(qn("a:latin"))
        for tag in ("a:ea", "a:cs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rpr, qn(tag))
                prev.addnext(el)
            el.set("typeface", FONT_NAME)
            prev = el
    if url:
        f.underline = True
        run.hyperlink.address = url


def fill_paragraph(p, runs, align=None, space_after=0, hang=None):
    if align is not None:
        p.alignment = align
    p.space_after = Pt(space_after)
    if hang is not None:
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", str(int(Inches(hang[0]))))
        ppr.set("indent", str(int(Inches(hang[1]))))
    for r in runs:
        run = p.add_run()
        run.text = r[0]
        style_run(run, r[1], r[2], r[3], r[4] if len(r) > 4 else None)


def add_text(shapes, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, name=None, wrap=True):
    tb = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        fill_paragraph(p, para["runs"], para.get("align"), para.get("space_after", 0), para.get("hang"))
    if name:
        tb.name = name
    return tb


def add_box(shapes, kind, x, y, w, h, fill=None, line=None, line_w=0.75, name=None, dash=False):
    shp = shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        apply_color(shp.fill.fore_color, fill)
    if line is None:
        shp.line.fill.background()
    else:
        apply_color(shp.line.color, line)
        shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shp.shadow.inherit = False
    if name:
        shp.name = name
    return shp


def add_label(shapes, rect, text, size, name, border=False):
    tb = add_text(shapes, *rect, [{"runs": [(text, size, RED, True)], "align": PP_ALIGN.CENTER}],
                  anchor=MSO_ANCHOR.MIDDLE, name=name, wrap=False)
    tb.fill.solid()
    apply_color(tb.fill.fore_color, C["white"])
    if border:
        tb.line.color.rgb = RED
        tb.line.width = Pt(1)
    return tb


def set_cell(cell, paras, fill, align):
    cell.fill.solid()
    apply_color(cell.fill.fore_color, fill)
    cell.margin_left = cell.margin_right = Inches(0.06)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    for i, runs in enumerate(paras):
        fill_paragraph(tf.paragraphs[0] if i == 0 else tf.add_paragraph(), runs, align)
    tcpr = cell._tc.get_or_add_tcPr()
    for side in "LRTB":
        for el in tcpr.findall(qn(f"a:ln{side}")):
            tcpr.remove(el)
    for i, side in enumerate("LRTB"):
        ln = etree.Element(qn(f"a:ln{side}"), w=str(int(Pt(0.75))), cap="flat", cmpd="sng", algn="ctr")
        if side in "TB":
            etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr")).set("val", str(RULE))
        else:
            etree.SubElement(ln, qn("a:noFill"))
        tcpr.insert(i, ln)


def add_table(shapes, x, y, col_w, rows, row_h, name):
    gf = shapes.add_table(len(rows), len(col_w), Inches(x), Inches(y), Inches(sum(col_w)), Inches(sum(row_h)))
    gf.name = name
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    tbl._tbl.tblPr.find(qn("a:tableStyleId")).text = NO_STYLE_TABLE
    for i, w in enumerate(col_w):
        tbl.columns[i].width = Inches(w)
    for r, (cells, fill) in enumerate(rows):
        tbl.rows[r].height = Inches(row_h[r])
        for c, (paras, align) in enumerate(cells):
            set_cell(tbl.cell(r, c), paras, fill, align)
    return gf


def cell_height(paras, width_in, pad=0.1):
    return sum(M.lines("".join(r[0] for r in runs), runs[0][1], (width_in - 0.12) * 72, runs[0][3])
               * runs[0][1] * LINE_FACTOR / 72 for runs in paras) + pad


def content_bounds(gray, scale):
    step = max(1, gray.width // 640)
    small = gray.reduce(step) if step > 1 else gray
    w, h = small.size
    data = small.tobytes()
    edge_l, edge_r = data[0::w], data[w - 1::w]

    # 全幅のバナーや罫線があっても余白と判定できるよう、背景色ではなく端の列との一致で見る
    def margin(x, edge):
        return sum(1 for a, b in zip(data[x::w], edge) if abs(a - b) <= 10) >= h * 0.985

    left = next((x for x in range(w) if not margin(x, edge_l)), 0)
    right = next((x for x in range(w - 1, -1, -1) if not margin(x, edge_r)), w - 1)
    if right - left < w * 0.4:
        return 0, gray.width
    pad = 16 * scale
    return max(0, left * step - pad), min(gray.width, (right + 1) * step + pad)


def ink_ratio(gray, region):
    x1, y1, x2, y2 = (int(round(v)) for v in region)
    x1, y1, x2, y2 = max(0, x1), max(0, y1), min(gray.width, x2), min(gray.height, y2)
    if x2 - x1 < 3 or y2 - y1 < 3:
        return 0.0
    patch = gray.crop((x1, y1, x2, y2))
    step = max(1, min(patch.size) // 24)
    if step > 1:
        patch = patch.reduce(step)
    pw, ph = patch.size
    px = patch.tobytes()
    bg = sorted(px)[len(px) // 2]
    on = [abs(v - bg) > 48 for v in px]
    # 罫線や枠線は文字ではないので、patch を横断する直線は数えない
    rows = {y for y in range(ph) if sum(on[y * pw:(y + 1) * pw]) > 0.8 * pw}
    cols = {x for x in range(pw) if sum(on[x::pw]) > 0.8 * ph}
    n = sum(1 for i, v in enumerate(on) if v and i // pw not in rows and i % pw not in cols)
    return n / (pw * ph)


def edge_ink(gray, px, x, y, w, h):
    iw, ih = gray.size

    def line(points):
        vals = [px[min(int(a), iw - 1), min(int(b), ih - 1)] for a, b in points]
        if not vals:
            return 0.0
        bg = sorted(vals)[len(vals) // 2]
        return sum(abs(v - bg) > 40 for v in vals) / len(vals)

    total = 0.0
    if x > 1:
        total += line((x, y + i) for i in range(0, int(h), 3))
    if x + w < iw - 1:
        total += line((x + w, y + i) for i in range(0, int(h), 3))
    if y > 1:
        total += line((x + i, y) for i in range(0, int(w), 3))
    if y + h < ih - 1:
        total += line((x + i, y + h) for i in range(0, int(w), 3))
    return total


def compute_crop(gray, scale, boxes, bounds, aspect, page_top):
    iw, ih = gray.size
    m = 40 * scale
    ux1 = min(b[0] for b in boxes)
    uy1 = min(b[1] for b in boxes)
    ux2 = max(b[0] + b[2] for b in boxes)
    uy2 = max(b[1] + b[3] for b in boxes)
    cx1 = max(0, min(bounds[0], ux1 - m))
    cx2 = min(iw, max(bounds[1], ux2 + m))
    avail = cx2 - cx1
    w = min(avail, max(ux2 - ux1 + 2 * m, MIN_CROP_CSS_W * scale))
    h = w / aspect
    if uy2 - uy1 + 2 * m > h:
        h = uy2 - uy1 + 2 * m
        w = min(avail, max(w, h * aspect))
    h = min(h, ih)
    x = cx1 if ux2 + m <= cx1 + w else max(cx1, min((ux1 + ux2) / 2 - w / 2, cx2 - w))
    y = 0 if page_top and uy2 + m <= h else max(0, min((uy1 + uy2) / 2 - h / 2, ih - h))

    # 枠だけから決めた切り出しは端で文字を切りやすいので、少し広げる/ずらす候補から端の文字量が少ないものを選ぶ。
    # 広げるほど文字が小さくなるのでコストに含め、上端は下げない（ページ見出しを落とさないため）
    px = gray.load()
    keep = 12 * scale

    def holds(a, b, c, d):
        return ux1 - keep >= a and uy1 - keep >= b and ux2 + keep <= a + c and uy2 + keep <= b + d

    base = edge_ink(gray, px, x, y, w, h)
    best = (round(base, 2), 0.0, x, y, w, h)
    for grow in (1.0, 1.05, 1.1):
        w2, h2 = min(w * grow, iw), min(h * grow, ih)
        for fx in (-0.06, -0.03, 0, 0.03, 0.06):
            for fy in (-0.06, -0.03, 0, 0.03, 0.06):
                x2 = min(max(x + (w - w2) / 2 + fx * w2, 0), iw - w2)
                y2 = min(max(y + (h - h2) / 2 + fy * h2, 0), ih - h2)
                if y2 > y + 0.5 or not holds(x2, y2, w2, h2):
                    continue
                cost = round(edge_ink(gray, px, x2, y2, w2, h2) + 1.5 * (grow - 1), 2)
                moved = abs(x2 - x) / w + abs(y2 - y) / h + (grow - 1)
                best = min(best, (cost, moved, x2, y2, w2, h2))
    return best[2:] if base - best[0] >= 0.04 else (x, y, w, h)


def overlap(a, b):
    w = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
    h = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    return w * h if w > 0 and h > 0 else 0.0


def inside(a, b):
    return (a[0] >= b[0] - 1e-6 and a[1] >= b[1] - 1e-6 and
            a[0] + a[2] <= b[0] + b[2] + 1e-6 and a[1] + a[3] <= b[1] + b[3] + 1e-6)


def distance(rect, cx, cy):
    dx = max(rect[0] - cx, 0, cx - rect[0] - rect[2])
    dy = max(rect[1] - cy, 0, cy - rect[1] - rect[3])
    return (dx * dx + dy * dy) ** 0.5


def separate_frames(frames, raws, line_w=2.25 / 72):
    # 隣接する要素の枠は境界を要素間の中央に置き、線同士に隙間を残す。
    # 要素間が狭いときは線幅ぶんだけ要素側へ食い込ませ、文字を横切るほどには縮めない
    for i in range(len(frames)):
        for j in range(i + 1, len(frames)):
            a, b = frames[i], frames[j]
            if overlap(a, b) <= 0 or overlap(raws[i], raws[j]) > 0:
                continue
            ra, rb = raws[i], raws[j]
            gap_x = max(ra[0], rb[0]) - min(ra[0] + ra[2], rb[0] + rb[2])
            gap_y = max(ra[1], rb[1]) - min(ra[1] + ra[3], rb[1] + rb[3])
            k = 0 if gap_x >= gap_y else 1
            (lo, rlo), (hi, rhi) = sorted(((a, ra), (b, rb)), key=lambda t: t[1][k])
            g = rhi[k] - (rlo[k] + rlo[k + 2])
            mid = rlo[k] + rlo[k + 2] + g / 2
            d = max(line_w * 0.75, min(0.04, g / 2))
            end = hi[k] + hi[k + 2]
            lo[k + 2] = min(lo[k] + lo[k + 2], mid - d) - lo[k]
            hi[k] = max(hi[k], mid + d)
            hi[k + 2] = end - hi[k]


def place_label(idx, frames, img_rect, placed, ink_at, size, prefer=None):
    fx, fy, fw, fh = frames[idx]
    g = 0.04
    lw, lh = size
    ly = fy if fh >= lh * 1.6 else fy + (fh - lh) / 2
    cands = [
        ("left", fx - lw - g, ly),
        ("top", fx, fy - lh - g),
        ("right", fx + fw + g, ly),
        ("top-center", fx + (fw - lw) / 2, fy - lh - g),
        ("bottom-center", fx + (fw - lw) / 2, fy + fh + g),
        ("left-bottom", fx - lw - g, fy + fh - lh),
        ("bottom", fx, fy + fh + g),
        ("top-right", fx + fw - lw, fy - lh - g),
        ("inside", fx + g, fy + g),
    ]
    margin_rect = (img_rect[0] - 0.4, img_rect[1] - 0.1, img_rect[2] + 0.4, img_rect[3] + 0.2)
    others = [f for i, f in enumerate(frames) if i != idx]
    best, best_cost = None, None
    for order, (key, x, y) in enumerate(cands):
        r = (x, y, lw, lh)
        cost = order * 0.02 + sum(overlap(r, o) for o in others + placed) * 100
        cx, cy = x + lw / 2, y + lh / 2
        own = distance(frames[idx], cx, cy)
        # 自分の枠より隣の枠に近い位置だと、どちらの枠の番号か読み違えられる
        if any(distance(o, cx, cy) < own * 0.9 for o in others):
            cost += 0.5
        if key == "inside":
            cost += 0.6
        # 文字を隠さないことを位置の好みより優先する（1%の差でも候補順を覆す重み）
        if inside(r, img_rect):
            cost += 0.1 + round(ink_at(r), 2) * 10
        elif inside(r, margin_rect):
            cost += 0.08 + round(ink_at(r), 2) * 10
        else:
            cost += 1000
        if key == prefer:
            cost -= 10
        if best_cost is None or cost < best_cost:
            best, best_cost = r, cost
    return best


def place_picture(slide, path, img_size, crop, area, name="画面キャプチャ"):
    iw, ih = img_size
    cx, cy, cw, ch = crop
    s = min(area[2] / cw, area[3] / ch)
    dw, dh = cw * s, ch * s
    x = area[0] + (area[2] - dw) / 2
    y = area[1]
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(dw), Inches(dh))
    pic.crop_left = cx / iw
    pic.crop_right = 1 - (cx + cw) / iw
    pic.crop_top = cy / ih
    pic.crop_bottom = 1 - (cy + ch) / ih
    pic.line.color.rgb = RULE
    pic.line.width = Pt(0.75)
    pic.name = name
    return (x, y, dw, dh), s


def frame_rect(box, crop, origin, s, pad_px, area):
    raw = [origin[0] + (box[0] - crop[0]) * s, origin[1] + (box[1] - crop[1]) * s, box[2] * s, box[3] * s]
    x, y = raw[0] - pad_px * s, raw[1] - pad_px * s
    w, h = raw[2] + 2 * pad_px * s, raw[3] + 2 * pad_px * s
    if w < MIN_FRAME[0]:
        x, w = x - (MIN_FRAME[0] - w) / 2, MIN_FRAME[0]
    if h < MIN_FRAME[1]:
        y, h = y - (MIN_FRAME[1] - h) / 2, MIN_FRAME[1]
    x1, y1 = max(x, area[0]), max(y, area[1])
    x2, y2 = min(x + w, area[0] + area[2]), min(y + h, area[1] + area[3])
    return [x1, y1, x2 - x1, y2 - y1], raw


def fit_crop(crop, aspect):
    cx, cy, cw, ch = crop
    if cw / ch < aspect:
        return cx, cy, cw, cw / aspect
    return cx, cy, ch * aspect, ch


class Geometry:
    def __init__(self, sw, sh, title=None, area=None, footer_top=None):
        self.sw, self.sh = sw, sh
        if area:
            self.mx, self.right = area[0], area[0] + area[2]
            self.title = title
            self.body_y = area[1]
            self.body_b = area[1] + area[3]
        elif title:
            self.mx, self.right = title[0], title[0] + title[2]
            self.title = title
            self.body_y = title[1] + title[3] + 0.18
            self.body_b = footer_top - 0.12 if footer_top else sh - 0.55
        else:
            self.mx, self.right = 0.45, sw - 0.45
            self.title = (0.45, 0.3, sw - 0.9, 0.56)
            self.body_y = 1.08
            self.body_b = sh - 0.55
        cw = self.right - self.mx
        self.panel_w = min(3.95, cw * 0.32)
        self.gap = 0.3
        self.img_w = cw - self.gap - self.panel_w
        self.cap_h = 0.32
        self.img_y = self.body_y + self.cap_h + 0.1
        self.img_h = self.body_b - self.img_y
        self.panel_x = self.mx + self.img_w + self.gap
        self.pad = 0.2


def theme_font(master, kind="minor"):
    theme = master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme.blob)
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    font = root.find(f".//{a}fontScheme/{a}{kind}Font")
    if font is None:
        return DEFAULT_FONT
    ea = font.find(f"{a}ea")
    jpan = next((f for f in font.findall(f"{a}font") if f.get("script") == "Jpan"), None)
    latin = font.find(f"{a}latin")
    for el in (ea, jpan, latin):
        if el is not None and el.get("typeface"):
            return el.get("typeface")
    return DEFAULT_FONT


def scheme_rgb(master, key):
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    name = master._element.find(qn("p:clrMap")).get(key, key)
    color = etree.fromstring(master.part.part_related_by(RT.THEME).blob).find(f".//{a}clrScheme/{a}{name}")[0]
    return RGBColor.from_string(color.get("lastClr") or color.get("val"))


def mix(base, color, ratio):
    return RGBColor(*(round(b + (c - b) * ratio) for b, c in zip(base, color)))


def hit_fill(background):
    # 取得行の薄赤は、差し込み先の背景色に赤を混ぜて作る（白背景なら HIT_BG と同じ色）。
    # 文字はテーマの文字色のままなので、ダークテーマでも読める。暗い背景では同じ割合だと見分けにくいので赤を強める
    return mix(background, RED, 0x11 / 0xFF if sum(background) > 3 * 0x80 else 0.25)


def title_size(master):
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    style = master._element.find(qn("p:txStyles") + "/" + qn("p:titleStyle"))
    rpr = style.find(f"{a}lvl1pPr/{a}defRPr") if style is not None else None
    return int(rpr.get("sz")) / 100 if rpr is not None and rpr.get("sz") else 32


def content_area(master):
    for lay in master.slide_layouts:
        phs = list(lay.placeholders)
        types = [ph.placeholder_format.type for ph in phs]
        bodies = [ph for ph in phs if ph.placeholder_format.type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)]
        if PP_PLACEHOLDER.TITLE in types and len(bodies) == 1 and len(types) - len(bodies) - 1 == sum(
                1 for t in types if t in FOOTER_TYPES):
            b = bodies[0]
            if b.left is not None and b.width is not None:
                return emu_in(b.left), emu_in(b.top), emu_in(b.width), emu_in(b.height)
    return None


def pick_layout(prs, ref_slide):
    masters = [ref_slide.slide_layout.slide_master] if ref_slide is not None else list(prs.slide_masters)
    best = None
    for master in masters:
        for lay in master.slide_layouts:
            types = [ph.placeholder_format.type for ph in lay.placeholders]
            if PP_PLACEHOLDER.TITLE not in types:
                continue
            extra = sum(1 for t in types if t != PP_PLACEHOLDER.TITLE and t not in FOOTER_TYPES)
            if best is None or extra < best[0]:
                best = (extra, lay)
    if best is None:
        raise SystemExit("差し込み先の資料にタイトル付きのレイアウトがありません")
    return best[1]


def emu_in(v):
    return Emu(v).inches


class Builder:
    def __init__(self, deck, base_dir, args):
        global M, FONT_NAME, C
        self.deck = deck
        self.base = base_dir
        self.max = args.max_per_slide
        self.layout = {i["no"]: i for i in deck.get("data_layout", [])}
        self.title = deck.get("title", "スクレイピング手順")
        self.warnings, self.overflows = [], []
        self.images = {}
        self.only = set(args.only.split(",")) if args.only else None
        self.inserting = args.base is not None
        self.prs = Presentation(args.base) if self.inserting else Presentation()
        # このスクリプトで作った資料（図形名「フッター」のテキストボックスがある）へ足すときは、
        # テーマに合わせると前後のスライドと書式がずれるため、元と同じ自前の書式で作る
        self.template = self.inserting and not any(
            shape.name == "フッター" for slide in self.prs.slides for shape in slide.shapes)
        if self.inserting:
            self.n_existing = len(self.prs.slides)
            self.insert_at = min(max(1, args.insert_at or self.n_existing + 1), self.n_existing + 1)
            ref = self.prs.slides[self.insert_at - 2] if self.insert_at > 1 else (
                self.prs.slides[0] if self.n_existing else None)
            self.content_layout = pick_layout(self.prs, ref)
            default_sections = "overview,steps,summary"
        else:
            self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
            self.n_existing, self.insert_at = 0, 1
            self.content_layout = self.prs.slide_layouts[5]
            default_sections = "cover,overview,steps,summary"
        if self.template:
            master = self.content_layout.slide_master
            tp = next(ph for ph in self.content_layout.placeholders
                      if ph.placeholder_format.type == PP_PLACEHOLDER.TITLE)
            feet = [emu_in(ph.top) for ph in self.content_layout.placeholders
                    if ph.placeholder_format.type in FOOTER_TYPES and ph.top is not None]
            title_rect = (emu_in(tp.left), emu_in(tp.top), emu_in(tp.width), emu_in(tp.height))
            self.geo = Geometry(emu_in(self.prs.slide_width), emu_in(self.prs.slide_height), title_rect,
                                content_area(master), min(feet) if feet else None)
            self.footer_src = next((s for s in list(self.prs.slides)[:self.n_existing]
                                    if any(ph.placeholder_format.type in FOOTER_TYPES for ph in s.placeholders)), None)
            FONT_NAME = None
            M = Metrics(theme_font(master))
            self.title_metrics = Metrics(theme_font(master, "major"))
            self.title_pt = title_size(master)
            bg, tx = scheme_rgb(master, "bg1"), scheme_rgb(master, "tx1")
            self.hit_bg = hit_fill(bg)
            # グレーアウトも差し込み先の背景色と文字色を混ぜて作る（白地に黒文字なら GRAY_BG / GRAY_TX と同じ色）
            self.gray_bg, self.gray_tx = mix(bg, tx, 0x26 / 0xFF), mix(bg, tx, 0x80 / 0xFF)
        else:
            self.geo = Geometry(13.333, 7.5)
            FONT_NAME = DEFAULT_FONT
            M = Metrics(DEFAULT_FONT)
            self.hit_bg, self.gray_bg, self.gray_tx = HIT_BG, GRAY_BG, GRAY_TX
        if self.only:
            default_sections = "steps"
        self.sections = set((args.sections or default_sections).split(","))
        C = PALETTES[self.template]
        self.slides = self.plan()

    def image(self, path, scale):
        key = str(path)
        if key not in self.images:
            if not path.exists():
                raise SystemExit(f"画像が見つかりません: {path}")
            gray = Image.open(path).convert("L")
            self.images[key] = {"path": path, "size": gray.size, "gray": gray,
                                "bounds": content_bounds(gray, scale)}
        return self.images[key]

    def url_of(self, shot):
        if shot.get("link") is False:
            return None
        return shot.get("url") or shot.get("captured_url")

    def plan(self):
        op_no = 0
        slides = []
        aspect = self.geo.img_w / self.geo.img_h
        for idx, shot in enumerate(self.deck["shots"]):
            sid = shot.get("id", idx + 1)
            steps = shot.get("steps", [])
            for st in steps:
                if st["kind"] == "op":
                    op_no += 1
                    st["_no"] = st.get("no", op_no)
                elif st["kind"] == "data":
                    if st["item"] not in self.layout:
                        raise SystemExit(f"データレイアウトに No.{st['item']} がありません（shot '{sid}'）")
                    st["_no"] = st["item"]
                elif st["kind"] != "mark":
                    raise SystemExit(f"shot '{sid}' に不明な kind があります: {st['kind']}")
            if self.only and sid not in self.only:
                continue
            if not shot.get("image"):
                raise SystemExit(f"shot '{sid}' に image がありません（先に capture.py を実行）")
            img = self.image((self.base / shot["image"]).resolve(), shot.get("scale", 1))
            iw, ih = img["size"]
            for st in steps:
                if st["kind"] == "op" and not st.get("box"):
                    self.warnings.append(f"shot '{sid}' の手順{circled(st['_no'])}に赤枠がありません（一覧にのみ表示）")
                if st["kind"] == "data" and not self.layout[st["_no"]].get("on_screen", True):
                    item = self.layout[st["_no"]]
                    self.warnings.append(f"画面外の項目 No.{item['no']}「{item['name']}」に赤枠が付いています")
                b = st.get("box")
                if b and (b[0] >= iw or b[1] >= ih or b[0] + b[2] <= 0 or b[1] + b[3] <= 0):
                    raise SystemExit(f"shot '{sid}' の赤枠が画像の外にあります: {b}（画像 {iw}x{ih}）")
            scale = shot.get("scale", 1)
            data = [s for s in steps if s["kind"] == "data"]
            ops = [s for s in steps if s["kind"] == "op"]
            marks = [s for s in steps if s["kind"] == "mark"]
            groups = []
            if data or (marks and not ops):
                groups.append(("data", data + marks))
            if ops:
                groups.append(("op", ops + ([] if data else marks)))
            mixed = len(groups) > 1
            for kind, group in groups:
                chunks = self.chunk(group, kind, scale)
                for n, chunk in enumerate(chunks, 1):
                    if kind == "op":
                        title = shot.get("title")
                    elif shot.get("data_title"):
                        title = shot["data_title"]
                    elif mixed or len(chunks) > 1 or not shot.get("title"):
                        title = self.items_title(chunk)
                    else:
                        title = shot["title"]
                    boxes = [s["box"] for s in chunk if s.get("box")]
                    crop = (compute_crop(img["gray"], scale, boxes, img["bounds"], aspect, shot.get("page_top", True))
                            if boxes else (0, 0, iw, min(ih, iw / aspect)))
                    slides.append({"shot": shot, "shot_index": idx, "kind": kind, "steps": chunk, "title": title,
                                   "part": (n, len(chunks)), "crop": crop, "img": img, "last": False})
            if slides and slides[-1]["shot_index"] == idx:
                slides[-1]["last"] = True
        return slides

    def items_title(self, steps):
        names = [self.layout[st["_no"]]["name"] for st in steps if st["kind"] == "data"]
        if not names:
            return "表示の違い"
        for n in range(len(names), 0, -1):
            text = "・".join(names[:n]) + ("など" if n < len(names) else "") + "を取得"
            if len(text) <= 26:
                return text
        return "取得項目"

    def chunk(self, steps, kind, scale):
        if kind == "data":
            steps = sorted(steps, key=lambda s: (s.get("box") or [0, 0])[1])
        # 縦に離れた枠を1枚に詰めると画像が縮んで読めなくなるため、離れていれば別スライドに分ける
        max_span = 900 * scale
        chunks, cur = [], []
        for st in steps:
            boxes = [s["box"] for s in cur + [st] if s.get("box")]
            too_tall = bool(cur) and bool(boxes) and \
                max(b[1] + b[3] for b in boxes) - min(b[1] for b in boxes) > max_span
            if cur and (len(cur) >= self.max or too_tall):
                chunks.append(cur)
                cur = []
            cur.append(st)
        if cur:
            chunks.append(cur)
        if kind == "data":
            chunks = [sorted(c, key=lambda s: (s["kind"] == "mark", s.get("_no", 0))) for c in chunks]
        return chunks

    def screens_of_items(self):
        where = {}
        for shot in self.deck["shots"]:
            for st in shot.get("steps", []):
                if st["kind"] == "data":
                    where.setdefault(st["_no"], {}).setdefault(shot["screen"], self.url_of(shot))
        return where

    def build(self, out_path):
        where = self.screens_of_items()
        for no, item in self.layout.items():
            if item.get("on_screen", True) and no not in where:
                self.warnings.append(f"データレイアウト No.{no}「{item['name']}」がどの画面にも対応付けられていません")
        if "cover" in self.sections and self.inserting:
            self.warnings.append("差し込みモードでは表紙を作りません（差し込み先の表紙を使ってください）")
        elif "cover" in self.sections:
            self.title_slide()
        if "overview" in self.sections:
            self.overview_slide()
        if "steps" in self.sections:
            for i, sl in enumerate(self.slides):
                self.content_slide(sl, i)
        if "summary" in self.sections:
            self.summary_slides(where)
        if self.inserting:
            self.move_new_slides()
        Path(out_path).parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(out_path)

    def move_new_slides(self):
        lst = self.prs.slides._sldIdLst
        ids = list(lst)
        new, old = ids[self.n_existing:], ids[:self.n_existing]
        pos = self.insert_at - 1
        for el in new:
            lst.remove(el)
        for i, el in enumerate(new):
            lst.insert(pos + i, el)
        ext = self.prs.part._element.find(qn("p:extLst"))
        if ext is None:
            return
        # セクションのある資料は全スライドがどこかのセクションに属していないと修復ダイアログが出る
        for e in ext:
            if e.get("uri") != SECTION_EXT:
                continue
            sections = e.findall(f".//{{{P14}}}section")
            anchor = old[pos - 1].get("id") if pos > 0 else (old[0].get("id") if old else None)
            target = None
            for sec in sections:
                ids_in = [x.get("id") for x in sec.iter(f"{{{P14}}}sldId")]
                if anchor in ids_in:
                    target = (sec, ids_in.index(anchor) + (1 if pos > 0 else 0))
                    break
            if target is None and sections:
                target = (sections[-1], len(list(sections[-1].iter(f"{{{P14}}}sldId"))))
            if target is None:
                return
            sec, at = target
            lst_el = sec.find(f"{{{P14}}}sldIdLst")
            for i, el in enumerate(new):
                node = etree.Element(f"{{{P14}}}sldId")
                node.set("id", el.get("id"))
                lst_el.insert(at + i, node)

    def new_slide(self):
        slide = self.prs.slides.add_slide(self.content_layout)
        for ph in list(slide.placeholders):
            if ph.placeholder_format.type != PP_PLACEHOLDER.TITLE:
                ph._element.getparent().remove(ph._element)
        return slide

    def finish_slide(self, slide):
        if not self.template:
            self.add_footer(slide)
        elif self.footer_src is not None:
            for ph in self.footer_src.placeholders:
                if ph.placeholder_format.type in FOOTER_TYPES:
                    el = copy.deepcopy(ph._element)
                    el.find(".//" + qn("p:cNvPr")).set("id", str(slide.shapes._next_shape_id))
                    slide.shapes._spTree.append(el)

    def add_footer(self, slide):
        g = self.geo
        add_text(slide.shapes, g.mx, g.sh - 0.42, 8.0, 0.25, [{"runs": [(self.title, 9, C["muted"], False)]}],
                 name="フッター")
        tb = add_text(slide.shapes, g.right - 1.0, g.sh - 0.42, 1.0, 0.25, [{"runs": [], "align": PP_ALIGN.RIGHT}],
                      name="スライド番号")
        fld = etree.SubElement(tb.text_frame.paragraphs[0]._p, qn("a:fld"))
        fld.set("id", "{" + str(uuid.uuid4()).upper() + "}")
        fld.set("type", "slidenum")
        rpr = etree.SubElement(fld, qn("a:rPr"), lang="ja-JP", sz="900")
        etree.SubElement(etree.SubElement(rpr, qn("a:solidFill")), qn("a:srgbClr")).set("val", str(C["muted"]))
        etree.SubElement(rpr, qn("a:latin"), typeface=FONT_NAME)
        etree.SubElement(rpr, qn("a:ea"), typeface=FONT_NAME)
        etree.SubElement(fld, qn("a:t")).text = str(len(self.prs.slides))

    def set_title(self, slide, text):
        t = slide.shapes.title
        tf = t.text_frame
        if self.template:
            # 資料側のタイトル書式を使い、1行に収まらないときだけ文字サイズを下げる
            size = self.title_pt
            while size > self.title_pt * 0.6 and self.title_metrics.width(text, size) / 72 > self.geo.title[2] - 0.2:
                size -= 1
            fill_paragraph(tf.paragraphs[0], [(text, size if size < self.title_pt else None, None, None)])
            return
        g = self.geo
        t.left, t.top, t.width, t.height = (Inches(v) for v in g.title)
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        size = 24
        while size > 16 and M.width(text, size, True) / 72 > g.title[2]:
            size -= 1
        fill_paragraph(tf.paragraphs[0], [(text, size, C["ink"], True)], PP_ALIGN.LEFT)

    def caption(self, slide, badge, badge_color, screen, url, variant=None):
        g = self.geo
        bw = M.width(badge, 11, True) / 72 + 0.3
        b = add_box(slide.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, g.mx, g.body_y + (g.cap_h - 0.28) / 2, bw, 0.28,
                    fill=badge_color, name="区分ラベル")
        b.adjustments[0] = 0.3
        tf = b.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        fill_paragraph(tf.paragraphs[0], [(badge, 11, C["white"], True)], PP_ALIGN.CENTER)
        avail = (g.img_w - bw - 0.14) * 72
        size = 14
        while size > 11 and M.width(screen, size, True) > avail * 0.6:
            size -= 1
        runs = [(screen, size, C["link"] if url else C["ink"], True, url)]
        if variant:
            rest = avail - M.width(screen, size, True)
            suffix = f"（表示パターン：{variant}）"
            ssize = 12
            while ssize > 9 and M.width(suffix, ssize) > rest:
                ssize -= 1
            while len(variant) > 1 and M.width(suffix, ssize) > rest:
                variant = variant[:-1]
                suffix = f"（表示パターン：{variant}…）"
            runs.append((suffix, ssize, C["muted"], False))
        add_text(slide.shapes, g.mx + bw + 0.14, g.body_y, g.img_w - bw - 0.14, g.cap_h, [{"runs": runs}],
                 anchor=MSO_ANCHOR.MIDDLE, name="画面名", wrap=False)

    def title_slide(self):
        d = self.deck
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[0])
        text_w = 7.3
        title, sub = slide.shapes.title, slide.placeholders[1]
        for ph, (y, h, anchor) in ((title, (1.55, 1.75, MSO_ANCHOR.BOTTOM)), (sub, (3.5, 0.9, MSO_ANCHOR.TOP))):
            ph.left, ph.top, ph.width, ph.height = Inches(0.8), Inches(y), Inches(text_w), Inches(h)
            tf = ph.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = anchor
        size = 34
        while size > 26 and M.width(self.title, size, True) / 72 > text_w:
            size -= 1
        lines = [self.title]
        if M.width(self.title, size, True) / 72 > text_w and " " in self.title:
            # 2行に折れるなら語の途中ではなく、中央に近い空白で改行する
            spaces = [i for i, c in enumerate(self.title) if c == " "]
            cut = min(spaces, key=lambda i: abs(i - len(self.title) / 2))
            lines = [self.title[:cut], self.title[cut + 1:]]
        tf = title.text_frame
        for i, line in enumerate(lines):
            fill_paragraph(tf.paragraphs[0] if i == 0 else tf.add_paragraph(), [(line, size, C["ink"], True)],
                           PP_ALIGN.LEFT)
        sub_text = d.get("subtitle", "")
        sub_size = 16
        while sub_size > 12 and M.width(sub_text, sub_size) / 72 > text_w:
            sub_size -= 1
        fill_paragraph(sub.text_frame.paragraphs[0], [(sub_text, sub_size, C["muted"], False)], PP_ALIGN.LEFT)
        n_all = len(self.layout)
        n_screen = sum(1 for i in self.layout.values() if i.get("on_screen", True))
        meta = []
        if d.get("site"):
            url = d.get("start_url") if d.get("link", True) else None
            site = [(d["site"], 12, C["link"] if url else C["ink"], False, url)]
            meta.append(("対象サイト", site))
        if n_all:
            meta.append(("取得項目", [(f"{n_all}項目（うち画面から取得 {n_screen}項目）", 12, C["ink"], False)]))
        meta.append(("作成日", [(d.get("date") or date.today().strftime("%Y/%m/%d"), 12, C["ink"], False)]))
        add_text(slide.shapes, 0.8, 4.7, text_w, 1.2,
                 [{"runs": [(f"{k}：", 12, C["muted"], False)] + v, "space_after": 4} for k, v in meta],
                 name="資料情報")
        if self.slides:
            first = self.slides[0]
            area = (8.5, 1.9, 4.35, 3.5)
            place_picture(slide, first["img"]["path"], first["img"]["size"],
                          fit_crop(first["crop"], area[2] / area[3]), area)

    def overview_slide(self):
        g = self.geo
        nodes = []
        for sl in self.slides:
            shot = sl["shot"]
            if not nodes or nodes[-1]["screen"] != shot["screen"]:
                nodes.append({"screen": shot["screen"], "slide": sl, "url": self.url_of(shot), "lines": []})
            line = f"表示パターン：{shot['variant']}" if shot.get("variant") else sl["title"]
            if line and line not in nodes[-1]["lines"] and sl["part"][0] == 1:
                nodes[-1]["lines"].append(line)
        n = len(nodes)
        if not n:
            return
        slide = self.new_slide()
        self.set_title(slide, "全体の流れ")
        per_row = n if n <= 4 else -(-n // 2)
        rows = -(-n // per_row)
        arrow = 0.42
        width = g.right - g.mx
        nw = (width - (per_row - 1) * arrow) / per_row
        captions = []
        for i, node in enumerate(nodes):
            paras = [{"runs": [(f"{i + 1}. ", 12, C["ink"], True),
                               (node["screen"], 12, C["link"] if node["url"] else C["ink"], True, node["url"])],
                      "space_after": 3}]
            for line in node["lines"]:
                paras.append({"runs": [(line, 11, C["muted"], False)], "space_after": 1})
            captions.append(paras)
        cap_h = max(self.paras_height(p, nw) for p in captions) + 0.05
        aspect = 1.25 if rows == 1 else 1.7
        th = nw / aspect
        legend_h = self.legend_height()
        legend_y = g.body_b - legend_h
        avail = legend_y - 0.25 - g.body_y
        if rows * (th + 0.12 + cap_h) + (rows - 1) * 0.2 > avail:
            th = max(0.6, (avail - (rows - 1) * 0.2) / rows - 0.12 - cap_h)
            aspect = nw / th
        block_h = rows * (th + 0.12 + cap_h) + (rows - 1) * 0.2
        if block_h > avail + 0.01:
            self.overflows.append("「全体の流れ」に画面が多すぎて収まりません。--max-per-slide を見直すか画面名を短くしてください")
        top = g.body_y + max(0.05, (avail - block_h) / 2)
        for i, (node, paras) in enumerate(zip(nodes, captions)):
            r, c = divmod(i, per_row)
            x = g.mx + c * (nw + arrow)
            y = top + r * (th + 0.12 + cap_h + 0.2)
            sl = node["slide"]
            place_picture(slide, sl["img"]["path"], sl["img"]["size"], fit_crop(sl["crop"], aspect), (x, y, nw, th))
            add_text(slide.shapes, x, y + th + 0.12, nw, cap_h, paras, name=f"画面{i + 1}")
            if c < per_row - 1 and i < n - 1:
                add_box(slide.shapes, MSO_SHAPE.RIGHT_ARROW, x + nw + 0.08, y + th / 2 - 0.13, arrow - 0.16, 0.26,
                        fill=RULE, name="矢印")
        self.legend(slide, legend_y, legend_h)
        self.finish_slide(slide)

    LEGEND = [("①", "画面遷移スライド", "赤枠の番号は操作の順番です"),
              ("③", "データ取得スライド", "赤枠の番号はデータレイアウトの項目No.です"),
              ("", "グレーの項目（クロール日時など）", "画面からは取得しない項目です")]

    def legend_height(self):
        g = self.geo
        tw = (g.right - g.mx - 2 * g.pad) / 3 - 1.05
        text_h = max(M.height(head, 12, tw, True) + 2 / 72 + M.height(desc, 11, tw) for _, head, desc in self.LEGEND)
        return 0.55 + max(text_h, 0.38) + 0.15

    def legend(self, slide, y, h):
        g = self.geo
        add_box(slide.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, g.mx, y, g.right - g.mx, h, fill=C["panel"],
                name="資料の見方").adjustments[0] = 0.08
        add_text(slide.shapes, g.mx + g.pad, y + 0.14, 3.0, 0.3, [{"runs": [("資料の見方", 12, C["muted"], True)]}])
        col_w = (g.right - g.mx - 2 * g.pad) / 3
        for i, (num, head, desc) in enumerate(self.LEGEND):
            x = g.mx + g.pad + i * col_w
            yy = y + 0.55
            if num:
                add_box(slide.shapes, MSO_SHAPE.RECTANGLE, x + 0.34, yy + 0.04, 0.46, 0.3, line=RED, line_w=2.25)
                add_label(slide.shapes, (x, yy + 0.04, LABEL, LABEL), num, 16, None)
            else:
                add_box(slide.shapes, MSO_SHAPE.RECTANGLE, x + 0.04, yy + 0.04, 0.76, 0.3, fill=self.gray_bg)
            add_text(slide.shapes, x + 0.95, yy, col_w - 1.05, h - 0.6,
                     [{"runs": [(head, 12, C["ink"], True)], "space_after": 2},
                      {"runs": [(desc, 11, C["muted"], False)]}])

    def content_slide(self, sl, index):
        g = self.geo
        shot = sl["shot"]
        slide = self.new_slide()
        part = f"（{sl['part'][0]}/{sl['part'][1]}）" if sl["part"][1] > 1 else ""
        self.set_title(slide, (sl.get("title") or shot["screen"]) + part)
        if sl["kind"] == "data":
            self.caption(slide, "データ取得", C["data"], shot["screen"], self.url_of(shot), shot.get("variant"))
        else:
            self.caption(slide, "画面遷移", C["nav"], shot["screen"], self.url_of(shot), shot.get("variant"))
        img = sl["img"]
        crop = sl["crop"]
        img_rect, s = place_picture(slide, img["path"], img["size"], crop, (g.mx, g.img_y, g.img_w, g.img_h))
        pad_px = 4 * shot.get("scale", 1)
        framed = [st for st in sl["steps"] if st.get("box")]
        pairs = [frame_rect(st["box"], crop, img_rect, s, pad_px, img_rect) for st in framed]
        frames = [f for f, _ in pairs]
        separate_frames(frames, [r for _, r in pairs])

        def ink_at(r):
            region = (crop[0] + (r[0] - img_rect[0]) / s, crop[1] + (r[1] - img_rect[1]) / s,
                      crop[0] + (r[0] + r[2] - img_rect[0]) / s, crop[1] + (r[1] + r[3] - img_rect[1]) / s)
            region = (max(region[0], crop[0]), max(region[1], crop[1]),
                      min(region[2], crop[0] + crop[2]), min(region[3], crop[1] + crop[3]))
            return ink_ratio(img["gray"], region)

        mark_no = 0
        labels = []
        for st in framed:
            if st["kind"] == "mark":
                mark_no += 1
                st["_mark"] = f"※{mark_no}"
                labels.append((st["_mark"], (0.42, LABEL), 13))
            else:
                labels.append((circled(st["_no"]), (LABEL, LABEL), 16))
        # 枠を先に全部描いてから番号を重ねる（後から描いた枠線が番号を横切らないように）
        for st, fr, (text, _, _) in zip(framed, frames, labels):
            desc = st.get("text") or (self.layout[st["_no"]]["name"] if st["kind"] == "data" else "")
            add_box(slide.shapes, MSO_SHAPE.RECTANGLE, *fr, line=RED, line_w=2.25 if st["kind"] != "mark" else 1.75,
                    dash=st["kind"] == "mark", name=f"赤枠{text} {desc[:20]}")
        placed = []
        for i, (st, (text, size, pt)) in enumerate(zip(framed, labels)):
            lab = place_label(i, frames, img_rect, placed, ink_at, size, st.get("label"))
            placed.append(lab)
            add_label(slide.shapes, lab, text, pt, f"番号{text}")
        for st in sl["steps"]:
            if st["kind"] == "mark" and "_mark" not in st:
                mark_no += 1
                st["_mark"] = f"※{mark_no}"
        add_box(slide.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, g.panel_x, g.img_y, g.panel_w, g.img_h, fill=C["panel"],
                name="右パネル").adjustments[0] = 0.04
        deferred = []
        if sl["kind"] == "data":
            deferred = self.data_panel(slide, sl, s)
        else:
            self.nav_panel(slide, sl, index)
        self.finish_slide(slide)
        if deferred:
            self.peek_slide(sl, deferred, s)

    def peek_slide(self, sl, peeks, scale_main):
        g = self.geo
        shot = sl["shot"]
        slide = self.new_slide()
        self.set_title(slide, (sl.get("title") or shot["screen"]) + "（表示が異なる場合）")
        self.caption(slide, "データ取得", C["data"], shot["screen"], self.url_of(shot), shot.get("variant"))
        width = g.right - g.mx
        img_w = width * 0.55
        row_h = (g.body_b - g.img_y - 0.2 * (len(peeks) - 1)) / len(peeks)
        y = g.img_y
        for pk, img, crop, s, pw, ph in self.peek_layout(peeks, img_w, scale_main * 1.5, row_h):
            rect, ps = place_picture(slide, img["path"], img["size"], crop, (g.mx, y, img_w, ph), name="表示差分の例")
            self.peek_mark(slide, pk, crop, rect, ps)
            runs = [(circled(pk["item"]) + " ", 14, RED, True)] if pk.get("item") else []
            runs.append((pk["text"], 13, C["ink"], False))
            paras = [{"runs": runs, "space_after": 6}]
            url = pk.get("url") or pk.get("captured_url")
            if url and pk.get("link", True):
                paras.append({"runs": [(pk.get("example", "表示例のページ"), 11, C["link"], False, url)]})
            add_text(slide.shapes, g.mx + img_w + 0.3, y, width - img_w - 0.3, ph, paras, name="表示差分")
            y += max(ph, 0.6) + 0.2
        self.finish_slide(slide)

    @staticmethod
    def peek_mark(slide, pk, crop, rect, s):
        if pk.get("mark_box"):
            mb = pk["mark_box"]
            fr = (rect[0] + (mb[0] - crop[0]) * s, rect[1] + (mb[1] - crop[1]) * s, mb[2] * s, mb[3] * s)
            add_box(slide.shapes, MSO_SHAPE.RECTANGLE, *fr, line=RED, line_w=1.5, dash=True, name="差分箇所")

    def mark_paras(self, steps, size):
        return [{"runs": [(st["_mark"] + " ", size, RED, True), (st.get("text", ""), size, C["ink"], False)],
                 "hang": (0.36, -0.36), "space_after": 4}
                for st in steps if st["kind"] == "mark"]

    def paras_height(self, paras, w):
        total = 0.0
        for p in paras:
            text = "".join(r[0] for r in p["runs"])
            size = max(r[1] for r in p["runs"])
            indent = p["hang"][0] if p.get("hang") else 0
            total += M.height(text, size, w - indent) + p.get("space_after", 0) / 72
        return total

    def nav_panel(self, slide, sl, index):
        g = self.geo
        shot = sl["shot"]
        x, w = g.panel_x + g.pad, g.panel_w - 2 * g.pad
        y = g.img_y + g.pad
        add_text(slide.shapes, x, y, w, 0.3, [{"runs": [("操作手順", 12, C["muted"], True)]}], name="見出し")
        y += 0.42
        hang = 0.34
        nxt = shot.get("next")
        if nxt is None and sl["part"][0] == sl["part"][1]:
            later = [o for o in self.slides[index + 1:]
                     if o["shot_index"] != sl["shot_index"] and not o["shot"].get("variant")]
            nxt = later[0]["shot"]["screen"] if later else None
        notes = shot.get("notes", []) if sl["last"] else []
        bottom = g.img_y + g.img_h - g.pad - (0.7 if nxt else 0)
        for size, dsize in ((14, 11), (13, 10), (12, 10)):
            paras = []
            for st in sl["steps"]:
                if st["kind"] != "op":
                    continue
                after = 2 if st.get("detail") else 8
                paras.append({"runs": [(circled(st["_no"]), size + 1, RED, True), ("\t" + st["text"], size, C["ink"], False)],
                              "hang": (hang, -hang), "space_after": after})
                if st.get("detail"):
                    paras.append({"runs": [(st["detail"], dsize, C["muted"], False)], "hang": (hang, 0), "space_after": 8})
            paras += self.mark_paras(sl["steps"], dsize)
            h = self.paras_height(paras, w)
            note_h = sum(M.height("※" + t, dsize, w) + 4 / 72 for t in notes)
            if y + h + (0.2 + note_h if notes else 0) <= bottom:
                break
        else:
            self.overflows.append(f"「{shot['screen']}」の操作手順が右パネルに収まりません。手順を分割するか文言を短くしてください")
        add_text(slide.shapes, x, y, w, h + 0.1, paras, name="操作手順")
        if notes:
            add_text(slide.shapes, x, y + h + 0.2, w, note_h + 0.05,
                     [{"runs": [("※" + t, dsize, C["muted"], False)], "space_after": 4} for t in notes], name="補足")
        if nxt:
            arrow = add_box(slide.shapes, MSO_SHAPE.PENTAGON, x, g.img_y + g.img_h - g.pad - 0.5, w, 0.5,
                            fill=C["nav"], name="次の画面")
            tf = arrow.text_frame
            tf.margin_left, tf.margin_right = Inches(0.15), Inches(0.3)
            tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.word_wrap = True
            fill_paragraph(tf.paragraphs[0], [("次の画面  ", 11, C["white"], False), (nxt, 12, C["white"], True)],
                           PP_ALIGN.LEFT)

    def peek_layout(self, peeks, w, scale_main, max_h):
        out = []
        for pk in peeks:
            path = (self.base / pk["image"]).resolve()
            img = self.image(path, pk.get("scale", 1))
            iw, ih = img["size"]
            crop = tuple(pk.get("box") or (0, 0, iw, ih))
            s = min(scale_main, w / crop[2], max_h / crop[3])
            out.append((pk, img, crop, s, crop[2] * s, crop[3] * s))
        return out

    def data_panel(self, slide, sl, scale_main):
        g = self.geo
        shot = sl["shot"]
        x, w = g.panel_x + g.pad, g.panel_w - 2 * g.pad
        y = g.img_y + g.pad
        bottom = g.img_y + g.img_h - g.pad
        hits = {st["_no"]: st for st in sl["steps"] if st["kind"] == "data"}
        here = {st["_no"] for o in self.slides if o["shot_index"] == sl["shot_index"] and o["kind"] == "data"
                for st in o["steps"] if st["kind"] == "data"}
        notes = shot.get("notes", []) if sl["last"] else []
        peeks = [pk for pk in shot.get("peeks", []) if pk.get("image") and (pk.get("item") in hits or (
            pk.get("item") is None and sl["part"][0] == sl["part"][1]))]
        diff = shot.get("diff", []) if sl["part"][0] == 1 else []
        if diff:
            diff_paras = [{"runs": [("通常表示との違い", 12, C["muted"], True)], "space_after": 4}] + \
                [{"runs": [("・" + t, 11, C["ink"], False)], "hang": (0.16, -0.16), "space_after": 3} for t in diff]
            dh = self.paras_height(diff_paras, w)
            add_text(slide.shapes, x, y, w, dh + 0.05, diff_paras, name="通常表示との違い")
            y += dh + 0.2
        add_text(slide.shapes, x, y, w, 0.3, [{"runs": [("取得項目（データレイアウト）", 12, C["muted"], True)]}],
                 name="見出し")
        y += 0.42
        items = sorted(self.layout.values(), key=lambda i: i["no"])
        candidates = [items,
                      [i for i in items if i["no"] in here or not i.get("on_screen", True)],
                      [i for i in items if i["no"] in hits]]
        mark_paras = self.mark_paras(sl["steps"], 11)

        def fit(with_peeks):
            attempt = None
            for subset in candidates:
                for size, peek_h in ((12, 1.3), (11, 1.1), (10, 0.9), (10, 0.7)):
                    col_w = self.panel_columns(subset, w, size)
                    rows, heights = self.panel_rows(subset, hits, shot, col_w, size)
                    gray = any(not i.get("on_screen", True) for i in subset)
                    laid = self.peek_layout(with_peeks, w, scale_main, peek_h)
                    extra = (0.12 + M.height(GRAY_NOTE, 10, w) if gray else 0) + \
                        (0.12 + self.paras_height(mark_paras, w) if mark_paras else 0) + \
                        sum(0.15 + M.height(circled(pk.get("item") or 0) + pk["text"], 11, w) + 0.06 + ph + 0.22
                            for pk, _, _, _, _, ph in laid) + (0.25 if laid else 0) + \
                        (0.12 + sum(M.height("※" + t, 10, w) + 0.06 for t in notes) if notes else 0)
                    attempt = (col_w, rows, heights, gray, laid)
                    if y + sum(heights) + extra <= bottom:
                        return attempt, True
            return attempt, False

        (col_w, rows, heights, gray, laid), ok = fit(peeks)
        deferred = []
        if not ok and peeks:
            # 表示差分の切り抜きが入りきらないときは、表を優先して切り抜きは直後の専用スライドへ回す
            (col_w, rows, heights, gray, laid), ok = fit([])
            deferred = peeks
        if not ok:
            self.overflows.append(f"「{shot['screen']}」の右パネルに収まりません。項目を分割してください")
        add_table(slide.shapes, x, y, col_w, rows, heights, "データレイアウト")
        y += sum(heights) + 0.12
        if gray:
            fh = M.height(GRAY_NOTE, 10, w)
            add_text(slide.shapes, x, y, w, fh, [{"runs": [(GRAY_NOTE, 10, C["muted"], False)]}], name="グレーアウト注記")
            y += fh + 0.12
        if mark_paras:
            mh = self.paras_height(mark_paras, w)
            add_text(slide.shapes, x, y, w, mh, mark_paras, name="注記")
            y += mh + 0.12
        if laid:
            y += 0.05
            add_text(slide.shapes, x, y, w, 0.25, [{"runs": [("表示が異なる場合", 11, C["muted"], True)]}],
                     name="表示差分見出し")
            y += 0.25
            for pk, img, crop, s, pw, ph in laid:
                runs = [(circled(pk["item"]) + " ", 12, RED, True)] if pk.get("item") else []
                runs.append((pk["text"], 11, C["ink"], False))
                th = M.height("".join(r[0] for r in runs), 11, w)
                add_text(slide.shapes, x, y, w, th, [{"runs": runs}], name="表示差分")
                y += th + 0.06
                rect, ps = place_picture(slide, img["path"], img["size"], crop, (x, y, pw, ph), name="表示差分の例")
                self.peek_mark(slide, pk, crop, rect, ps)
                url = pk.get("url") or pk.get("captured_url")
                if url and pk.get("link", True):
                    add_text(slide.shapes, x, y + ph + 0.03, w, 0.2,
                             [{"runs": [(pk.get("example", "表示例のページ"), 9, C["link"], False, url)]}],
                             name="表示差分の出典")
                y += ph + 0.28
        if notes:
            add_text(slide.shapes, x, y, w, max(0.3, bottom - y),
                     [{"runs": [("※" + t, 10, C["muted"], False)], "space_after": 4} for t in notes], name="補足")
        return deferred

    @staticmethod
    def panel_columns(items, w, size):
        type_w = max([M.width(i.get("type", ""), size) for i in items] + [M.width("型", size, True)]) / 72 + 0.16
        type_w = min(max(type_w, 0.7), 1.5)
        return [0.45, w - 0.45 - type_w, type_w]

    def panel_rows(self, items, hits, shot, col_w, size):
        rows = [([([[("No.", size, C["white"], True)]], PP_ALIGN.CENTER),
                  ([[("項目名", size, C["white"], True)]], PP_ALIGN.LEFT),
                  ([[("型", size, C["white"], True)]], PP_ALIGN.LEFT)], C["head"])]
        heights = [size * LINE_FACTOR / 72 + 0.1]
        where = self.screens_of_items()
        for it in items:
            no = it["no"]
            if not it.get("on_screen", True):
                fill, color, num_color, bold, sub = self.gray_bg, self.gray_tx, self.gray_tx, False, None
            elif no in hits:
                fill, color, num_color, bold, sub = self.hit_bg, C["ink"], RED, True, hits[no].get("detail")
            else:
                other = [s for s in where.get(no, {}) if s != shot["screen"]]
                fill, color, num_color, bold = C["white"], C["muted"], C["muted"], False
                sub = f"{other[0]}で取得" if other else None
            name = [[(it["name"], size, color, bold)]]
            if sub:
                name.append([(sub, size - 2, self.gray_tx if color is self.gray_tx else C["muted"], False)])
            cells = [([[(circled(no), size + 1, num_color, True)]], PP_ALIGN.CENTER),
                     (name, PP_ALIGN.LEFT),
                     ([[(it.get("type", ""), size, color, False)]], PP_ALIGN.LEFT)]
            heights.append(max(cell_height(name, col_w[1]), cell_height(cells[2][0], col_w[2])))
            rows.append((cells, fill))
        return rows, heights

    def summary_slides(self, where):
        g = self.geo
        items = sorted(self.layout.values(), key=lambda i: i["no"])
        if not items:
            return
        width = g.right - g.mx
        col_w = [w * width / 12.43 for w in (0.8, 2.9, 1.5, 3.4, 3.83)]
        size = 13 if width > 11 else 11
        y0 = g.body_y + 0.1
        note_w = min(8.0, width)
        note_h = M.height(GRAY_NOTE, 11, note_w)
        pages, cur, used = [], [], 0.42
        gray = False
        for it in items:
            off = not it.get("on_screen", True)
            color = self.gray_tx if off else C["ink"]
            if off:
                src = [[(it.get("origin") or "画面外", size, color, False)]]
            elif where.get(it["no"]):
                runs = []
                for k, (screen, url) in enumerate(where[it["no"]].items()):
                    if k:
                        runs.append(("、", size, color, False))
                    runs.append((screen, size, C["link"] if url else color, False, url))
                src = [runs]
            else:
                src = [[("（未対応付け）", size, color, False)]]
            cells = [([[(circled(it["no"]), size + 1, self.gray_tx if off else RED, True)]], PP_ALIGN.CENTER),
                     ([[(it["name"], size, color, False)]], PP_ALIGN.LEFT),
                     ([[(it.get("type", ""), size, color, False)]], PP_ALIGN.LEFT),
                     (src, PP_ALIGN.LEFT),
                     ([[(it.get("note", ""), size, color, False)]], PP_ALIGN.LEFT)]
            h = max(cell_height(cells[k][0], col_w[k], 0.14) for k in range(1, 5))
            if y0 + 0.42 + h + (0.15 + note_h if off else 0) > g.body_b:
                self.overflows.append(f"取得項目 No.{it['no']}「{it['name']}」の対応表の行が1ページに収まりません。"
                                      "取得元や備考を短くしてください")
            if cur and y0 + used + h + (0.15 + note_h if gray or off else 0) > g.body_b:
                pages.append(cur)
                cur, used = [], 0.42
                gray = False
            cur.append((cells, self.gray_bg if off else C["white"], h, off))
            used += h
            gray = gray or off
        pages.append(cur)
        for pi, chunk in enumerate(pages, 1):
            slide = self.new_slide()
            suffix = f"（{pi}/{len(pages)}）" if len(pages) > 1 else ""
            self.set_title(slide, "取得項目と取得元の対応" + suffix)
            rows = [([([[(t, size, C["white"], True)]], PP_ALIGN.CENTER if t == "No." else PP_ALIGN.LEFT)
                      for t in ("No.", "項目名", "型", "取得元", "備考")], C["head"])]
            heights = [0.42]
            for cells, fill, h, _ in chunk:
                heights.append(h)
                rows.append((cells, fill))
            add_table(slide.shapes, g.mx, y0, col_w, rows, heights, "取得項目一覧")
            if any(off for _, _, _, off in chunk):
                add_text(slide.shapes, g.mx, y0 + sum(heights) + 0.15, note_w, note_h,
                         [{"runs": [(GRAY_NOTE, 11, C["muted"], False)]}], name="グレーアウト注記")
            self.finish_slide(slide)


def merge_captures(deck, deck_dir, captures_dir):
    captured = json.loads((captures_dir / "deck.json").read_text(encoding="utf-8"))
    by_id = {s.get("id"): s for s in captured["shots"]}
    for shot in deck["shots"]:
        if shot.get("image"):
            shot["image"] = str((deck_dir / shot["image"]).resolve())
            for pk in shot.get("peeks", []):
                if pk.get("image"):
                    pk["image"] = str((deck_dir / pk["image"]).resolve())
            continue
        src = by_id.get(shot.get("id"))
        if src is None or not src.get("image"):
            continue
        if any(shot.get(key) != src.get(key) for key in
               ("goto", "setup", "wait_ms", "wait_for", "wait_url", "timeout_ms", "blur", "scroll_until", "capture_mode")):
            raise SystemExit(f"shot '{shot.get('id')}' の撮影定義が撮影時と違います。capture.py で撮り直してください")
        steps, src_steps = shot.get("steps", []), src.get("steps", [])
        if [s.get("target") for s in steps] != [s.get("target") for s in src_steps]:
            raise SystemExit(f"shot '{shot.get('id')}' の手順の数か target が撮影時と違います。capture.py で撮り直してください")
        shot["image"] = str((captures_dir / src["image"]).resolve())
        for key in ("image_size", "scale", "page_top", "captured_url"):
            if key in src:
                shot[key] = src[key]
        for st, sst in zip(steps, src_steps):
            if sst.get("box"):
                st["box"] = sst["box"]
        src_peeks = src.get("peeks", [])
        for i, pk in enumerate(shot.get("peeks", [])):
            if pk.get("image"):
                pk["image"] = str((deck_dir / pk["image"]).resolve())
            else:
                if i >= len(src_peeks) or not src_peeks[i].get("image") or any(
                        pk.get(key) != src_peeks[i].get(key)
                        for key in ("goto", "target", "mark", "pad", "setup", "wait_ms", "wait_for", "wait_url",
                                    "timeout_ms", "scroll_until", "capture_mode")):
                    raise SystemExit(f"shot '{shot.get('id')}' の peeks[{i}] の撮影定義が撮影時と違うか、画像がありません。"
                                     "capture.py で撮り直してください")
                for key in ("image", "image_size", "scale", "mark_box", "captured_url"):
                    if key in src_peeks[i]:
                        pk[key] = src_peeks[i][key]
                pk["image"] = str((captures_dir / src_peeks[i]["image"]).resolve())
    return deck


def main():
    # Windowsのパイプ出力は既定でcp932になり、WARNやエラーの日本語が文字化けするため
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="シナリオ（＋撮影結果）から、赤枠と番号を編集できるPPTX手順書を生成する")
    ap.add_argument("deck", type=Path, help="scenario.json（--captures と併用）または capture.py が出力した deck.json")
    ap.add_argument("-o", "--out", type=Path)
    ap.add_argument("--captures", type=Path, help="capture.py の出力先。文言は scenario.json、画像と座標はここから使う")
    ap.add_argument("--base", type=Path, help="差し込み先の既存pptx。そのデザイン（レイアウト・テーマ色・フォント・フッター）に合わせる")
    ap.add_argument("--insert-at", type=int, help="--base の何枚目の位置に差し込むか（1始まり。省略時は末尾）")
    ap.add_argument("--sections", help="作るスライド: cover,overview,steps,summary のカンマ区切り")
    ap.add_argument("--only", help="指定した shot id のスライドだけ作る（カンマ区切り）。番号は全体の通し番号のまま")
    ap.add_argument("--max-per-slide", type=int, default=5, help="1スライドに載せる赤枠の上限")
    args = ap.parse_args()
    deck = json.loads(args.deck.read_text(encoding="utf-8"))
    if args.captures:
        deck = merge_captures(deck, args.deck.parent, args.captures)
    out = args.out or args.deck.with_name(re.sub(r'[\\/:*?"<>|\s]+', "_", deck.get("title", "deck")) + ".pptx")
    if args.base and args.base.resolve() == out.resolve():
        raise SystemExit("--base と出力先が同じです。元の資料を上書きしないよう別名で出力してください")
    b = Builder(deck, args.deck.parent, args)
    b.build(out)
    for w in b.warnings:
        print("WARN:", w, file=sys.stderr)
    for w in b.overflows:
        print("NG:", w, file=sys.stderr)
    print(f"pptx: {out}  ({len(b.prs.slides)} slides)")
    sys.exit(3 if b.overflows else 0)


if __name__ == "__main__":
    main()
